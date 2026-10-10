"""Read any deck — .pptx/.potx (also .ppt/.odp via LibreOffice) or a PDF export — into an inventory an agent can reason
about and edit against: every shape with its id, name, box (1280-wide canvas px), effective text style, image quality,
table/chart data and notes; plus real renders, a contact sheet and mechanical issues (overflow, low-res, small text…)."""
import colorsys
import hashlib
import pathlib
import re
import shutil
import subprocess
import tempfile

from pptx.oxml.ns import qn

from .util import CANVAS_W, contrast, hex6, log, read_json, write_json

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
DEFAULT_PT = 18.0
CLRMAP_DEFAULT = {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}
MONO = re.compile(r"mono|consolas|courier|menlo", re.I)


# ---------- colour helpers ----------
def _apply_mods(hexv, el):
    """Approximate DrawingML colour transforms (lumMod/lumOff/tint/shade) on an sRGB hex."""
    r, g, b = (int(hexv[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    for m in el:
        tag, val = m.tag.split("}")[-1], m.get("val")
        if val is None:
            continue
        v = int(val) / 100000
        if tag == "lumMod":
            l *= v
        elif tag == "lumOff":
            l += v
        elif tag == "tint":
            l = l + (1 - l) * (1 - v)
        elif tag == "shade":
            l *= v
    r, g, b = colorsys.hls_to_rgb(h, max(0.0, min(1.0, l)), s)
    return "".join(f"{round(c * 255):02X}" for c in (r, g, b))


class Theme:
    def __init__(self, prs):
        master = prs.slide_masters[0]
        from lxml import etree
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        root = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
        self.colors, self.fonts = {}, {}
        cs = root.find(f".//{A}clrScheme")
        if cs is not None:
            for el in cs:
                slot = el.tag.split("}")[-1]
                if len(el):
                    c = el[0]
                    v = c.get("val") if c.tag == A + "srgbClr" else c.get("lastClr")
                    if v:
                        self.colors[slot] = hex6(v)
        for tag, key in (("majorFont", "major"), ("minorFont", "minor")):
            latin = root.find(f".//{A}fontScheme/{A}{tag}/{A}latin")
            if latin is not None and latin.get("typeface"):
                self.fonts[key] = latin.get("typeface")
        cm = master._element.find(qn("p:clrMap"))
        self.clrmap = dict(CLRMAP_DEFAULT)
        if cm is not None:
            self.clrmap.update({k: v for k, v in cm.attrib.items()})

    def color(self, fill_parent):
        """Hex of the first colour element under a solidFill-bearing element (or the colour element itself)."""
        if fill_parent is None:
            return None
        c = fill_parent if fill_parent.tag in (A + "srgbClr", A + "schemeClr", A + "sysClr", A + "prstClr") else None
        if c is None:
            sf = fill_parent.find(A + "solidFill")
            if sf is None or not len(sf):
                return None
            c = sf[0]
        tag = c.tag.split("}")[-1]
        if tag == "srgbClr":
            base = hex6(c.get("val"))
        elif tag == "sysClr":
            base = hex6(c.get("lastClr") or "000000")
        elif tag == "schemeClr":
            slot = self.clrmap.get(c.get("val"), c.get("val"))
            base = self.colors.get(slot)
            if not base:
                return None
        else:
            return None
        return _apply_mods(base, c) if len(c) else base

    def typeface(self, tf):
        if not tf:
            return None
        if tf.startswith("+mj"):
            return self.fonts.get("major")
        if tf.startswith("+mn"):
            return self.fonts.get("minor")
        return tf


# ---------- style inheritance ----------
TITLE_TYPES = {"title", "ctrTitle"}
BODY_TYPES = {"body", "subTitle", "obj", None}


def _ph_info(sp):
    ph = sp.find(f".//{P}nvPr/{P}ph")
    if ph is None:
        return None
    return {"type": ph.get("type") or "body" if ph.get("idx") else (ph.get("type") or "obj"), "idx": ph.get("idx")}


def _find_ph(tree_root, info):
    """Placeholder on a layout/master matching a slide placeholder (idx first, then type)."""
    if tree_root is None or info is None:
        return None
    cands = tree_root.findall(f".//{P}sp")
    if info.get("idx") is not None:
        for sp in cands:
            ph = sp.find(f".//{P}nvPr/{P}ph")
            if ph is not None and ph.get("idx") == info["idx"]:
                return sp
    want = info.get("type")
    alias = {"ctrTitle": "title", "subTitle": "body", "obj": "body"}
    for t in (want, alias.get(want)):
        for sp in cands:
            ph = sp.find(f".//{P}nvPr/{P}ph")
            if ph is not None and (ph.get("type") or "obj") == t:
                return sp
    return None


def _lvl_props(container, lvl):
    """defRPr + pPr of list level `lvl` in an a:lstStyle / txStyles child."""
    if container is None:
        return None, None
    p = container.find(f"{A}lvl{lvl + 1}pPr")
    if p is None:
        return None, None
    return p.find(A + "defRPr"), p


class StyleChain:
    """Ordered sources of default run/paragraph properties for one shape (most specific first)."""

    def __init__(self, sp, slide, prs_el, theme):
        self.theme = theme
        self.sources = []  # list of callables lvl -> (defRPr, pPr)
        self.title = False
        txBody = sp.find(P + "txBody")
        own = txBody.find(A + "lstStyle") if txBody is not None else None
        self.sources.append(lambda lvl, c=own: _lvl_props(c, lvl))
        info = _ph_info(sp)
        if info:
            self.title = info["type"] in TITLE_TYPES
            lay = _find_ph(slide.slide_layout._element, info)
            mas = _find_ph(slide.slide_layout.slide_master._element, info)
            for ph in (lay, mas):
                if ph is not None:
                    tb = ph.find(P + "txBody")
                    ls = tb.find(A + "lstStyle") if tb is not None else None
                    self.sources.append(lambda lvl, c=ls: _lvl_props(c, lvl))
            tx = slide.slide_layout.slide_master._element.find(qn("p:txStyles"))
            if tx is not None:
                key = "titleStyle" if self.title else ("bodyStyle" if info["type"] in BODY_TYPES else "otherStyle")
                st = tx.find(P + key)
                self.sources.append(lambda lvl, c=st: _lvl_props(c, lvl))
        dts = prs_el.find(qn("p:defaultTextStyle"))
        self.sources.append(lambda lvl, c=dts: _lvl_props(c, lvl))

    def resolve(self, rPr, lvl):
        out = {"sz": None, "b": None, "i": None, "latin": None, "color": None, "explicit_color": False}
        chain = [rPr] + [src(lvl)[0] for src in self.sources]
        for el in chain:
            if el is None:
                continue
            if out["sz"] is None and el.get("sz"):
                out["sz"] = int(el.get("sz")) / 100
            if out["b"] is None and el.get("b") is not None:
                out["b"] = el.get("b") in ("1", "true")
            if out["i"] is None and el.get("i") is not None:
                out["i"] = el.get("i") in ("1", "true")
            if out["latin"] is None:
                lat = el.find(A + "latin")
                if lat is not None and lat.get("typeface"):
                    out["latin"] = lat.get("typeface")
            if out["color"] is None:
                c = self.theme.color(el)
                if c:
                    out["color"] = c
                    out["explicit_color"] = el is rPr
        return out

    def align(self, pPr, lvl):
        if pPr is not None and pPr.get("algn"):
            return pPr.get("algn")
        for src in self.sources:
            p = src(lvl)[1]
            if p is not None and p.get("algn"):
                return p.get("algn")
        return "l"


# ---------- shapes ----------
def _xfrm_of(el):
    for path in (f"{P}spPr/{A}xfrm", f"{P}grpSpPr/{A}xfrm", f"{P}xfrm"):
        x = el.find(path)
        if x is not None:
            return x
    return None


def _shape_kind(shape):
    if shape._element.tag == P + "grpSp":
        return "group"
    if shape._element.tag == P + "pic":
        el = shape._element
        if el.find(f".//{P}nvPr/{A}videoFile") is not None or el.find(f".//{P}nvPr/{A}audioFile") is not None:
            return "media"
        return "image"
    if getattr(shape, "has_table", False) and shape.has_table:
        return "table"
    if getattr(shape, "has_chart", False) and shape.has_chart:
        return "chart"
    if shape._element.tag == P + "graphicFrame":
        return "other"
    if shape._element.tag == P + "cxnSp":
        return "line"
    if shape.has_text_frame and shape.text_frame.text.strip():
        return "text"
    return "shape"


def _image_info(shape, box, k):
    el = shape._element
    blip = el.find(f".//{A}blip")
    info = {"alt": el.find(f".//{P}cNvPr").get("descr") or ""}
    if blip is None:
        return info
    rid = blip.get(R + "embed")
    try:
        part = shape.part.related_part(rid)
    except KeyError:
        return info
    info["content_type"] = part.content_type
    info["partname"] = str(part.partname)
    info["bytes"] = len(part.blob)
    svg = el.find(".//{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip")
    info["svg"] = svg is not None
    try:
        from io import BytesIO

        from PIL import Image
        with Image.open(BytesIO(part.blob)) as im:
            nw, nh = im.size
    except Exception:  # noqa: BLE001 - emf/wmf/svg
        return info
    src = el.find(f".//{A}srcRect")
    crop = {s: (int(src.get(s, 0)) / 100000 if src is not None else 0.0) for s in ("l", "t", "r", "b")}
    vis_w = nw * max(0.01, 1 - crop["l"] - crop["r"])
    vis_h = nh * max(0.01, 1 - crop["t"] - crop["b"])
    info.update({"natural": [nw, nh], "crop": {s: round(v, 4) for s, v in crop.items() if v}})
    if box["w"] > 1 and box["h"] > 1:
        info["px_per_canvas_px"] = round(min(vis_w / box["w"], vis_h / box["h"]), 2)
        info["stretch"] = round((vis_w / vis_h) / (box["w"] / box["h"]), 3)
    return info


def _table_info(shape):
    tbl = shape.table
    rows = []
    for r in tbl.rows:
        rows.append([c.text_frame.text if c.text_frame is not None else "" for c in r.cells])
    return {"rows": len(tbl.rows), "cols": len(tbl.columns), "cells": rows}


def _chart_info(shape):
    ch = shape.chart
    out = {"type": str(ch.chart_type).split(".")[-1].split(" ")[0].lower(), "series": []}
    try:
        plot = ch.plots[0]
        out["categories"] = [str(c) for c in plot.categories]
        for s in plot.series:
            out["series"].append({"name": s.name, "values": [v for v in s.values]})
    except Exception as e:  # noqa: BLE001
        out["error"] = str(e)
    try:
        out["title"] = ch.chart_title.text_frame.text if ch.has_title else ""
    except Exception:  # noqa: BLE001
        pass
    return out


def _text_info(shape, chain, theme, k):
    tf = shape.text_frame
    body = tf._txBody
    bp = body.find(A + "bodyPr")
    autofit = None
    if bp is not None:
        if bp.find(A + "normAutofit") is not None:
            autofit = "shrink"
            scale = bp.find(A + "normAutofit").get("fontScale")
            if scale:
                autofit = f"shrink {int(scale) / 1000:.0f}%"
        elif bp.find(A + "spAutoFit") is not None:
            autofit = "resize-shape"
    paras = []
    for p in body.findall(A + "p"):
        pPr = p.find(A + "pPr")
        lvl = int(pPr.get("lvl", 0)) if pPr is not None else 0
        runs = []
        for r in p:
            tag = r.tag.split("}")[-1]
            if tag == "br":
                runs.append({"br": True})
                continue
            if tag not in ("r", "fld"):
                continue
            t = r.find(A + "t")
            text = t.text if t is not None and t.text else ""
            if not text:
                continue
            st = chain.resolve(r.find(A + "rPr"), lvl)
            pt = st["sz"] or DEFAULT_PT
            fam = theme.typeface(st["latin"]) or theme.fonts.get("major" if chain.title else "minor") or "Calibri"
            rPr = r.find(A + "rPr")
            caps = rPr is not None and rPr.get("cap") == "all"
            spc = int(rPr.get("spc", 0)) / 100 if rPr is not None and rPr.get("spc") else 0
            runs.append({"text": text, "pt": round(pt, 1), "size": round(pt / 0.75 * k, 2), "family": fam,
                         "weight": 700 if st["b"] else 400, "italic": bool(st["i"]), "hex": st["color"] or theme.colors.get("dk1", "000000"),
                         "color_known": st["color"] is not None, "caps": "all" if caps else "none", "spacing": spc})
        if not any(not x.get("br") for x in runs):
            continue
        bullet = pPr is not None and (pPr.find(A + "buChar") is not None or pPr.find(A + "buAutoNum") is not None)
        algn = chain.align(pPr, lvl)
        paras.append({"level": lvl, "align": {"l": "l", "ctr": "ctr", "r": "r", "just": "just"}.get(algn, "l"),
                      "bullet": bullet, "runs": runs})
    return {"paragraphs": paras, "autofit": autofit, "wrap": (bp.get("wrap") != "none") if bp is not None else True}


def walk_shapes(shapes, slide, ctx, transform=None, depth=0):
    """Flatten-aware walk: returns elements (groups keep children) with absolute canvas-px boxes."""
    k = ctx["k"]
    out = []
    for shape in shapes:
        el = shape._element
        try:
            x, y, w, h = (shape.left or 0), (shape.top or 0), (shape.width or 0), (shape.height or 0)
        except Exception:  # noqa: BLE001
            x = y = w = h = 0
        if transform:
            x, y, w, h = transform(x, y, w, h)
        box = {"x": round(x / 9525 * k, 1), "y": round(y / 9525 * k, 1), "w": round(w / 9525 * k, 1), "h": round(h / 9525 * k, 1)}
        kind = _shape_kind(shape)
        item = {"id": shape.shape_id, "name": shape.name, "kind": kind, "box": box}
        if getattr(shape, "rotation", 0):
            item["rotation"] = round(shape.rotation, 1)
        info = _ph_info(el)
        if info:
            item["placeholder"] = info["type"]
        if kind == "group":
            xf = _xfrm_of(el)
            if xf is not None and xf.find(A + "chExt") is not None:
                off, ext = xf.find(A + "off"), xf.find(A + "ext")
                choff, chext = xf.find(A + "chOff"), xf.find(A + "chExt")
                ox, oy, ex, ey = int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))
                cx0, cy0 = int(choff.get("x")), int(choff.get("y"))
                sx, sy = ex / max(1, int(chext.get("cx"))), ey / max(1, int(chext.get("cy")))

                def tf(a, b, c, d, ox=ox, oy=oy, cx0=cx0, cy0=cy0, sx=sx, sy=sy, parent=transform):
                    nx, ny, nw, nh = ox + (a - cx0) * sx, oy + (b - cy0) * sy, c * sx, d * sy
                    return parent(nx, ny, nw, nh) if parent else (nx, ny, nw, nh)
                item["children"] = walk_shapes(shape.shapes, slide, ctx, tf, depth + 1)
            else:
                item["children"] = walk_shapes(shape.shapes, slide, ctx, transform, depth + 1)
        elif kind == "image":
            item["image"] = _image_info(shape, box, k)
            if ctx.get("media_dir"):
                _dump_media(shape, item, ctx)
        elif kind == "table":
            item["table"] = _table_info(shape)
        elif kind == "chart":
            item["chart"] = _chart_info(shape)
        elif kind == "other":
            gd = el.find(f".//{A}graphicData")
            item["uri"] = gd.get("uri") if gd is not None else ""
            item["other"] = ("smartart" if "diagram" in item["uri"] else "ole" if "ole" in item["uri"] else "graphic")
        if kind in ("text", "shape") and shape.has_text_frame:
            chain = StyleChain(el, slide, ctx["prs_el"], ctx["theme"])
            t = _text_info(shape, chain, ctx["theme"], k)
            if t["paragraphs"]:
                item["text"] = t
        if kind in ("text", "shape"):
            spPr = el.find(P + "spPr")
            fill = ctx["theme"].color(spPr) if spPr is not None else None
            if fill:
                item["fill"] = fill
            elif spPr is not None and spPr.find(A + "blipFill") is not None:
                item["fill"] = "image"
            elif spPr is not None and spPr.find(A + "gradFill") is not None:
                item["fill"] = "gradient"
        out.append(item)
    return out


def _dump_media(shape, item, ctx):
    blip = shape._element.find(f".//{A}blip")
    if blip is None:
        return
    try:
        part = shape.part.related_part(blip.get(R + "embed"))
    except KeyError:
        return
    ext = str(part.partname).rsplit(".", 1)[-1]
    out = ctx["media_dir"] / f"slide-{ctx['slide_index']:02d}-id{shape.shape_id}.{ext}"
    out.write_bytes(part.blob)
    item["image"]["file"] = str(out)


def _walk(elements):
    for e in elements:
        yield e
        yield from _walk(e.get("children", []))


def _plain(item):
    t = item.get("text")
    if not t:
        return ""
    return "\n".join("".join(r["text"] if not r.get("br") else "\n" for r in p["runs"]) for p in t["paragraphs"])


def _transition(sld):
    for el in sld.iter():
        tag = el.tag.split("}")[-1]
        if tag == "transition":
            kids = [c.tag.split("}")[-1] for c in el if not c.tag.endswith("sndAc") and not c.tag.endswith("extLst")]
            return kids[0] if kids else "fade-like"
    return None


def _background(sld, theme):
    bg = sld.find(f"{P}cSld/{P}bg")
    if bg is None:
        return "layout"
    bgPr = bg.find(P + "bgPr")
    if bgPr is not None:
        if bgPr.find(A + "blipFill") is not None:
            return "image"
        if bgPr.find(A + "gradFill") is not None:
            return "gradient"
        c = theme.color(bgPr)
        return f"#{c}" if c else "solid"
    ref = bg.find(P + "bgRef")
    if ref is not None and len(ref):
        c = theme.color(ref[0])
        return f"#{c}" if c else "theme"
    return "layout"


def _title_of(elements):
    for e in _walk(elements):
        if e.get("placeholder") in ("title", "ctrTitle") and _plain(e).strip():
            return _plain(e).strip().replace("\n", " ")
    best = None
    cands = [e for e in _walk(elements) if e.get("text") and e["box"]["y"] < 400]
    worded = [e for e in cands if len(_plain(e).split()) >= 3]  # a big stat ("25") is not the title
    for e in worded or cands:
        if e:
            size = max(r["size"] for p in e["text"]["paragraphs"] for r in p["runs"] if not r.get("br"))
            if best is None or size > best[0] + 0.5 or (abs(size - best[0]) <= 0.5 and e["box"]["y"] < best[1]["box"]["y"]):
                best = (size, e)
    return _plain(best[1]).strip().replace("\n", " ") if best else ""


# ---------- inputs ----------
def convert_to_pptx(path, out_dir):
    from .render import find_soffice
    soffice = find_soffice()
    if not soffice:
        raise RuntimeError(f"{path.suffix} needs LibreOffice to convert to .pptx; save it as .pptx in PowerPoint instead")
    with tempfile.TemporaryDirectory() as tmp:
        profile = pathlib.Path(tmp, "profile").as_uri()
        subprocess.run([soffice, f"-env:UserInstallation={profile}", "--headless", "--norestore", "--convert-to", "pptx",
                        "--outdir", str(out_dir), str(path)], capture_output=True, timeout=300, check=True)
    out = out_dir / (path.stem + ".pptx")
    if not out.exists():
        raise RuntimeError(f"LibreOffice could not convert {path.name}")
    return out


def _sha(path):
    return hashlib.sha256(pathlib.Path(path).read_bytes()).hexdigest()


def read_pptx(path, media_dir=None):
    from .compile import open_template
    prs = open_template(path)
    theme = Theme(prs)
    sw, sh = prs.slide_width, prs.slide_height
    k = CANVAS_W * 9525 / sw
    ctx = {"k": k, "theme": theme, "prs_el": prs.part._element, "media_dir": media_dir}
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        ctx["slide_index"] = i
        sld = slide._element
        els = walk_shapes(slide.shapes, slide, ctx)
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame is not None else ""
        anims = len(sld.findall(f".//{P}timing//{P}cTn[@presetClass]"))
        words = sum(len(_plain(e).split()) for e in _walk(els)) + sum(
            len(c.split()) for e in _walk(els) if e.get("table") for row in e["table"]["cells"] for c in row)
        bg = _background(sld, theme)
        slides.append({"index": i, "slide_id": slide.slide_id, "layout": slide.slide_layout.name, "hidden": sld.get("show") == "0",
                       "title": _title_of(els), "background": bg, "transition": _transition(sld), "animations": anims,
                       "words": words, "notes": notes, "elements": els})
    layouts = {}
    for s in slides:
        layouts[s["layout"]] = layouts.get(s["layout"], 0) + 1
    deck = {"slide_size_in": [round(sw / 914400, 3), round(sh / 914400, 3)], "aspect": round(sw / sh, 3),
            "canvas": [CANVAS_W, round(CANVAS_W * sh / sw, 1)], "scale_px_per_pt": round(k / 0.75, 4),
            "theme": {"colors": theme.colors, "fonts": theme.fonts}, "layouts_used": layouts,
            "masters": len(prs.slide_masters), "sections": _sections(prs)}
    return deck, slides


def _sections(prs):
    out = []
    ids = {s.slide_id: i for i, s in enumerate(prs.slides, 1)}
    for sec in prs.part._element.iter("{http://schemas.microsoft.com/office/powerpoint/2010/main}section"):
        members = [ids.get(int(x.get("id"))) for x in sec.iter("{http://schemas.microsoft.com/office/powerpoint/2010/main}sldId")]
        out.append({"name": sec.get("name"), "slides": [m for m in members if m]})
    return out


def read_pdf(path, render_dir):
    import pymupdf as fitz
    doc = fitz.open(path)
    slides = []
    w0 = None
    for i, page in enumerate(doc, 1):
        pw, ph = page.rect.width, page.rect.height
        w0 = w0 or (pw, ph)
        k = CANVAS_W / pw
        if render_dir:
            page.get_pixmap(matrix=fitz.Matrix(1920 / pw, 1920 / pw), alpha=False).save(str(render_dir / f"slide-{i:02d}.png"))
        els = []
        n = 0
        for b in page.get_text("dict")["blocks"]:
            n += 1
            x0, y0, x1, y1 = b["bbox"]
            box = {"x": round(x0 * k, 1), "y": round(y0 * k, 1), "w": round((x1 - x0) * k, 1), "h": round((y1 - y0) * k, 1)}
            if b.get("type") == 1:
                nw, nh = b.get("width", 0), b.get("height", 0)
                els.append({"id": n, "name": f"image {n}", "kind": "image", "box": box,
                            "image": {"natural": [nw, nh], "px_per_canvas_px": round(min(nw / max(1, box["w"]), nh / max(1, box["h"])), 2)}})
                continue
            runs, widths = [], []
            for ln in b.get("lines", []):
                widths.append(ln["bbox"][2] - ln["bbox"][0])
                for sp in ln.get("spans", []):
                    if not sp["text"].strip():
                        continue
                    font = sp["font"]
                    weight = 700 if (sp["flags"] & 16 or re.search(r"bold|black|heavy|extrabold", font, re.I)) else 400
                    fam = re.sub(r"^[A-Z]{6}\+", "", font).split("-")[0].split(",")[0]
                    runs.append({"text": sp["text"] + " ", "pt": round(sp["size"], 1), "size": round(sp["size"] * k, 2), "family": fam,
                                 "weight": weight, "italic": bool(sp["flags"] & 2), "hex": f"{sp['color']:06X}", "color_known": True,
                                 "caps": "none", "spacing": 0})
            if not runs:
                continue
            runs[-1]["text"] = runs[-1]["text"].rstrip()
            ratio = round(widths[-1] / max(widths), 2) if len(widths) > 1 and max(widths) else 1
            els.append({"id": n, "name": f"text {n}", "kind": "text", "box": box,
                        "text": {"paragraphs": [{"level": 0, "align": "l", "bullet": False, "lines": len(widths),
                                                 "last_line_ratio": ratio, "runs": runs}]}})
        words = sum(len(_plain(e).split()) for e in els)
        slides.append({"index": i, "slide_id": i, "layout": "pdf-page", "hidden": False, "title": _title_of(els), "background": "pdf",
                       "transition": None, "animations": 0, "words": words, "notes": "", "elements": els})
    deck = {"slide_size_in": [round(w0[0] / 72, 3), round(w0[1] / 72, 3)] if w0 else None,
            "aspect": round(w0[0] / w0[1], 3) if w0 else None, "canvas": [CANVAS_W, round(CANVAS_W * w0[1] / w0[0], 1)] if w0 else None,
            "theme": {}, "layouts_used": {}, "masters": 0, "sections": [], "pdf": True}
    doc.close()
    return deck, slides


# ---------- checks ----------
def _issues(deck, slides, render_info, render_dir, scale):
    from PIL import Image

    from .fonts import FontIndex
    from .qa import _bg_color
    issues = []

    def add(slide, code, severity, msg, el=None):
        issues.append({"slide": slide, "code": code, "severity": severity, "msg": msg,
                       "el": f"#{el['id']} {el['name']}" if el else ""})

    fams = {}
    for s in slides:
        for e in _walk(s["elements"]):
            for p in (e.get("text") or {}).get("paragraphs", []):
                for r in p["runs"]:
                    if not r.get("br"):
                        fams[r["family"]] = fams.get(r["family"], 0) + 1
    if not deck.get("pdf"):
        idx = FontIndex()
        deck["fonts_used"] = {f: {"runs": n, "installed": idx.has(f)} for f, n in sorted(fams.items(), key=lambda x: -x[1])}
        for f, info in deck["fonts_used"].items():
            if not info["installed"]:
                add(0, "font-missing", "error", f"'{f}' is not installed here: renders (and edits) will show a substitute font. "
                                                f"Install it (`deck.py fonts fetch \"{f}\" --install` if it is a Google font) before judging the design")
    else:
        deck["fonts_used"] = {f: {"runs": n} for f, n in sorted(fams.items(), key=lambda x: -x[1])}
    rs = {s["index"]: s for s in (render_info or {}).get("slides", [])}
    for s in slides:
        i = s["index"]
        if s["hidden"]:
            add(i, "hidden", "info", "hidden slide (skipped in the slideshow)")
        if s["words"] > 90:
            add(i, "dense", "warn", f"{s['words']} words on one slide")
        img = None
        ren = render_dir / f"slide-{i:02d}.png" if render_dir else None
        if ren and ren.exists():
            img = Image.open(ren)
        by_id = {e["id"]: e for e in _walk(s["elements"])}
        for o in (rs.get(i) or {}).get("overflow", []):
            # PowerPoint's text bounds include descender/line-spacing slack of a few pt; real overflow is part of a line
            e = by_id.get(o["id"]) or {}
            pts = [r["pt"] for p in (e.get("text") or {}).get("paragraphs", []) for r in p["runs"] if not r.get("br")]
            if o["over_y"] > max(4.0, 0.6 * max(pts or [DEFAULT_PT])) or o["over_x"] > 4.0:
                add(i, "overflow", "error", f"text runs {max(o['over_y'], o['over_x']):.0f}pt past its box in PowerPoint",
                    {"id": o["id"], "name": o["name"]})
        H = deck["canvas"][1] if deck.get("canvas") else 720
        for e in _walk(s["elements"]):
            b = e["box"]
            if e["kind"] in ("text", "table", "chart") and (b["x"] < -2 or b["y"] < -2 or b["x"] + b["w"] > CANVAS_W + 2 or b["y"] + b["h"] > H + 2):
                add(i, "off-slide", "warn", "content extends past the slide edge", e)
            if e.get("placeholder") and e["kind"] == "shape" and not e.get("text") and e.get("placeholder") not in ("dt", "ftr", "sldNum", "pic"):
                add(i, "empty-placeholder", "info", "empty placeholder (shows 'Click to add…' while editing)", e)
            im = e.get("image") or {}
            if im.get("px_per_canvas_px") is not None and im["px_per_canvas_px"] < 1.0:
                add(i, "low-res", "warn", f"image is upscaled {1 / max(0.01, im['px_per_canvas_px']):.1f}× on a 1280-px canvas — soft when projected", e)
            if im.get("stretch") and abs(im["stretch"] - 1) > 0.03:
                add(i, "stretched", "error", f"image aspect distorted by {abs(im['stretch'] - 1) * 100:.0f}%", e)
            if e["kind"] == "image" and not im.get("alt") and not deck.get("pdf"):
                add(i, "alt-text", "info", "picture has no alt text", e)
            runs = [r for p in (e.get("text") or {}).get("paragraphs", []) for r in p["runs"] if not r.get("br")]
            if runs and e.get("placeholder") not in ("dt", "ftr", "sldNum"):
                small = min(r["pt"] for r in runs)
                if small < 10:
                    add(i, "small-text", "warn" if small >= 8 else "error", f"{small:.0f}pt text is hard to read when presented", e)
                r0 = next((r for r in runs if r["color_known"]), None)
                if img is not None and r0 and e["kind"] == "text":
                    bg = _bg_color(img, b, scale, r0["hex"])
                    if bg:
                        cr = contrast(r0["hex"], bg)
                        large = r0["size"] >= 24 or (r0["size"] >= 18.66 and r0["weight"] >= 700)
                        need = 3.0 if large else 4.5
                        if cr < need:
                            add(i, "contrast", "warn", f"contrast {cr:.2f}:1 (#{r0['hex']} on ~#{bg}) below WCAG AA {need}:1", e)
            if e["kind"] == "other":
                add(i, "non-native", "info", f"{e.get('other')} object: read it from the render; edit it in PowerPoint", e)
    return issues


# ---------- outline ----------
def _style_tag(e):
    runs = [r for p in e["text"]["paragraphs"] for r in p["runs"] if not r.get("br")]
    r = max(runs, key=lambda r: len(r["text"]))
    w = " bold" if r["weight"] >= 700 else ""
    return f"{r['pt']:g}pt {r['family']}{w} #{r['hex']}"


def _outline_el(e, depth, lines):
    pad = "  " * depth
    b = e["box"]
    where = f"@{b['x']:.0f},{b['y']:.0f} {b['w']:.0f}×{b['h']:.0f}"
    head = f"{pad}- #{e['id']} \"{e['name']}\" {e['kind']}" + (f" ({e['placeholder']})" if e.get("placeholder") else "") + f" {where}"
    if e.get("text"):
        lines.append(f"{head} [{_style_tag(e)}" + (f", autofit {e['text']['autofit']}" if e["text"].get("autofit") else "") + "]")
        for p in e["text"]["paragraphs"]:
            t = "".join(r["text"] if not r.get("br") else " ⏎ " for r in p["runs"])
            lines.append(f"{pad}    {'  ' * p['level']}{'• ' if p['bullet'] else ''}{t}")
    elif e["kind"] == "image":
        im = e.get("image") or {}
        nat = "×".join(str(v) for v in im.get("natural", [])) or "?"
        extra = f", file {pathlib.Path(im['file']).name}" if im.get("file") else ""
        lines.append(f"{head} {nat}px{' svg' if im.get('svg') else ''}" + (f", alt: {im['alt'][:80]}" if im.get("alt") else "") + extra)
    elif e["kind"] == "table":
        t = e["table"]
        lines.append(f"{head} {t['rows']}×{t['cols']}")
        for row in t["cells"][:8]:
            lines.append(f"{pad}    | " + " | ".join(c.replace("\n", " ")[:40] for c in row) + " |")
        if t["rows"] > 8:
            lines.append(f"{pad}    | … {t['rows'] - 8} more rows")
    elif e["kind"] == "chart":
        c = e["chart"]
        lines.append(f"{head} {c.get('type')} chart" + (f" “{c['title']}”" if c.get("title") else ""))
        if c.get("categories"):
            lines.append(f"{pad}    categories: {', '.join(c['categories'][:12])}")
        for s in c.get("series", [])[:6]:
            lines.append(f"{pad}    {s['name']}: {', '.join(f'{v:g}' if isinstance(v, (int, float)) else str(v) for v in s['values'][:12])}")
    elif e["kind"] == "group":
        lines.append(f"{head} ({len(e.get('children', []))} items)")
        for ch in e.get("children", []):
            _outline_el(ch, depth + 1, lines)
    else:
        extra = f" fill #{e['fill']}" if e.get("fill") and e["fill"] not in ("image", "gradient") else (f" fill {e['fill']}" if e.get("fill") else "")
        lines.append(f"{head}{extra}" + (f" {e.get('other')}" if e.get("other") else ""))


def outline(src, deck, slides, issues):
    t = deck.get("theme") or {}
    lines = [f"# {src.name} — {len(slides)} slides", ""]
    if deck.get("slide_size_in"):
        lines.append(f"Size {deck['slide_size_in'][0]}×{deck['slide_size_in'][1]} in (aspect {deck['aspect']}). "
                     f"Boxes are canvas px on a {deck['canvas'][0]:.0f}×{deck['canvas'][1]:.0f} canvas; text sizes in pt.")
    if t.get("fonts"):
        lines.append(f"Theme fonts: heading {t['fonts'].get('major')}, body {t['fonts'].get('minor')}")
    if t.get("colors"):
        lines.append("Theme colours: " + " ".join(f"{k}=#{v}" for k, v in t["colors"].items() if k not in ("hlink", "folHlink")))
    if deck.get("fonts_used"):
        lines.append("Fonts used: " + ", ".join(f"{f} ({v['runs']} runs{'' if v.get('installed', True) else ', NOT INSTALLED'})"
                                                for f, v in deck["fonts_used"].items()))
    if deck.get("pipeline"):
        lines.append(f"Built by Slidekick: {deck['pipeline']}")
    if deck.get("sections"):
        lines.append("Sections: " + "; ".join(f"{s['name']} ({', '.join(map(str, s['slides'])) or 'empty'})" for s in deck["sections"]))
    lines.append("")
    lines.append("## Storyline (titles only)")
    for s in slides:
        lines.append(f"{s['index']}. {s['title'] or '(no title)'}" + (" [hidden]" if s["hidden"] else ""))
    for s in slides:
        meta = [f"layout: {s['layout']}", f"bg: {s['background']}", f"{s['words']} words"]
        if s["transition"]:
            meta.append(f"transition: {s['transition']}")
        if s["animations"]:
            meta.append(f"{s['animations']} builds")
        if s["hidden"]:
            meta.append("HIDDEN")
        lines += ["", f"## {s['index']} · {s['title'] or '(no title)'}", f"_{' · '.join(meta)}_", ""]
        for e in s["elements"]:
            _outline_el(e, 0, lines)
        if s["notes"].strip():
            lines += ["", "> Notes: " + s["notes"].strip().replace("\n", "\n> ")]
        si = [x for x in issues if x["slide"] == s["index"] and x["severity"] != "info"]
        if si:
            lines += [""] + [f"⚠ {x['code']}: {x['msg']}" + (f" — {x['el']}" if x["el"] else "") for x in si]
    deck_issues = [x for x in issues if not x["slide"]]
    if deck_issues:
        lines += ["", "## Deck-level issues"] + [f"- **{x['severity']}** `{x['code']}` {x['msg']}" for x in deck_issues]
    return "\n".join(lines) + "\n"


# ---------- entry ----------
def inspect_deck(src, out_dir, render=True, engine="auto", media=False):
    from .qa import contact_sheet
    from .render import render as do_render
    src = pathlib.Path(src).resolve()
    out_dir = pathlib.Path(out_dir).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    ext = src.suffix.lower()
    pptx = src
    if ext in (".ppt", ".odp", ".pps", ".ppsx"):
        pptx = convert_to_pptx(src, out_dir)
        log(f"converted {src.name} -> {pptx.name}")
    render_dir = out_dir / "render"
    media_dir = None
    if media:
        media_dir = out_dir / "media"
        media_dir.mkdir(exist_ok=True)
    render_info = None
    if ext == ".pdf":
        if render_dir.exists():
            shutil.rmtree(render_dir)
        render_dir.mkdir(parents=True)
        deck, slides = read_pdf(src, render_dir)
        render_info = {"engine": "pymupdf", "slides": []}
    else:
        deck, slides = read_pptx(pptx, media_dir)
        cj = src.parent / "build" / "compile.json"
        if cj.exists():
            try:
                built = read_json(cj).get("sha256")
                if built:
                    deck["pipeline"] = "unchanged since the last build" if built == _sha(src) else \
                        "MODIFIED after the last build (someone edited the .pptx; rebuilding from the canvas would discard those edits)"
            except Exception:  # noqa: BLE001
                pass
        if render:
            try:
                render_info = do_render(pptx, render_dir, engine=engine, text_bounds=True)
            except Exception as e:  # noqa: BLE001 - inventory is still useful without pictures
                log(f"render failed: {e}")
                deck["render_error"] = str(e)
    scale = 1920 / CANVAS_W
    issues = _issues(deck, slides, render_info, render_dir if (render_dir.exists() and (render or ext == ".pdf")) else None, scale)
    sheets = []
    pngs = sorted(render_dir.glob("slide-*.png")) if render_dir.exists() else []
    for n in range(0, len(pngs), 15):
        chunk = pngs[n:n + 15]
        name = "contact-sheet.png" if len(pngs) <= 15 else f"contact-sheet-{n // 15 + 1:02d}.png"
        contact_sheet([(str(int(p.stem.split("-")[1])), p) for p in chunk], out_dir / name)
        sheets.append(str(out_dir / name))
    inv = {"source": str(src), "file": str(pptx), "sha256": _sha(src), "renderer": (render_info or {}).get("engine"),
           "deck": deck, "slides": slides, "issues": issues, "contact_sheets": sheets}
    write_json(out_dir / "inventory.json", inv)
    (out_dir / "outline.md").write_text(outline(src, deck, slides, issues), encoding="utf-8")
    return inv
