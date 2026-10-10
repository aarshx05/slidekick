"""Code generation: layout.json (Low-IR) -> native, fully editable .pptx."""
import hashlib
import pathlib

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from . import calib
from . import ooxml as X
from .fonts import GENERIC, GENERIC_MAP, FontIndex, embeddable, stack_families, system_font_dir
from .util import CANVAS_W, hex6, log, read_json, write_json

NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED, "stacked-column": XL_CHART_TYPE.COLUMN_STACKED,
    "percent-column": XL_CHART_TYPE.COLUMN_STACKED_100, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "stacked-bar": XL_CHART_TYPE.BAR_STACKED, "percent-bar": XL_CHART_TYPE.BAR_STACKED_100,
    "line": XL_CHART_TYPE.LINE, "line-markers": XL_CHART_TYPE.LINE_MARKERS, "area": XL_CHART_TYPE.AREA,
    "stacked-area": XL_CHART_TYPE.AREA_STACKED, "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "scatter": XL_CHART_TYPE.XY_SCATTER,
}


class Fx:
    """Per-slide conversion context handed to the XML builders."""

    def __init__(self, ctx, slide):
        self.ctx, self.slide = ctx, slide
        self.k = ctx.k
        self.alpha = 1.0
        self.theme = ctx.theme_colors
        self.lang = ctx.lang

    def emu(self, px):
        return int(round(px * 9525 * self.k))

    def pt(self, px):
        return px * 0.75 * self.k

    def font(self, run):
        return self.ctx.resolve_font(run)

    def theme_font(self, typeface):
        if self.ctx.theme_fonts.get("major") and typeface == self.ctx.theme_fonts["major"]:
            return "+mj-lt"
        if self.ctx.theme_fonts.get("minor") and typeface == self.ctx.theme_fonts["minor"]:
            return "+mn-lt"
        return typeface

    def link_rid(self, url):
        return self.slide.part.relate_to(url, RT.HYPERLINK, is_external=True)


class Ctx:
    def __init__(self, prs, k, brand, media_dir, lang="en-US", brand_dir=None):
        self.prs, self.k, self.brand, self.media_dir, self.lang = prs, k, brand or {}, media_dir, lang
        base = pathlib.Path(brand_dir or ".")
        brand_fonts = [str(base / p) for p in (self.brand.get("fonts", {}).get("files") or []) if (base / p).exists()]
        self.index = FontIndex(extra_files=brand_fonts)
        self.font_cache, self.missing, self.used_faces = {}, {}, {}
        self.svg_parts = {}
        self.theme_colors, self.theme_fonts = read_theme(prs)
        self.warnings = []
        self.baselines = {}

    def style_of(self, p):
        run = next((r for r in p["runs"] if not r.get("br")), p["runs"][0])
        typeface, bold, italic = self.resolve_font(run)
        sz = max(100, min(400000, int(round(run["size"] * 0.75 * self.k * 100))))
        spc = max(0, int(round(p["line_height"] * 0.75 * self.k * 100)))
        return typeface, bold, italic, sz, spc

    def baseline(self, p):
        st = self.style_of(p)
        o = self.baselines.get(calib.key(*st)) or calib.model(st[4])
        return {a: o[a] / self.k for a in ("t", "ctr", "b")}

    def resolve_font(self, run):
        fam = run.get("platform_font")
        if not fam:
            css = stack_families(run.get("family", "Arial"))
            fam = next((f for f in css if f.lower() not in GENERIC), None) or GENERIC_MAP.get((css or ["sans-serif"])[0].lower(), "Arial")
        key = (fam, run.get("weight", 400), run.get("italic", False))
        if key not in self.font_cache:
            typeface, bold, italic, face = self.index.resolve(fam, run.get("weight", 400), run.get("italic", False))
            if face is None:
                self.missing.setdefault(fam, 0)
                self.missing[fam] += 1
            else:
                self.used_faces[(typeface, bold, italic)] = face
            self.font_cache[key] = (typeface, bold, italic)
        return self.font_cache[key]


def read_theme(prs):
    master = prs.slide_masters[0]
    root = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    colors = {}
    cs = root.find(f".//{A}clrScheme")
    if cs is not None:
        for slot in ["accent1", "accent2", "accent3", "accent4", "accent5", "accent6", "dk1", "lt1", "dk2", "lt2"]:
            el = cs.find(A + slot)
            if el is None or not len(el):
                continue
            c = el[0]
            val = c.get("val") if c.tag == A + "srgbClr" else c.get("lastClr")
            if val and hex6(val) not in colors:
                colors[hex6(val)] = slot
    fonts = {}
    for tag, key in (("majorFont", "major"), ("minorFont", "minor")):
        latin = root.find(f".//{A}fontScheme/{A}{tag}/{A}latin")
        if latin is not None and latin.get("typeface"):
            fonts[key] = latin.get("typeface")
    return colors, fonts


# ---------- helpers ----------
def _set_name(shape, name, descr=None):
    if name:
        shape.name = name[:250]
    if descr:
        shape._element.xpath("./*[1]/p:cNvPr")[0].set("descr", descr[:1000])


def _remove_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def prepare_raster(path, box_w, box_h, crop, media_dir):
    from PIL import Image, ImageOps

    path = pathlib.Path(path)
    ok_ext = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"}
    try:
        im = Image.open(path)
    except Exception:  # noqa: BLE001
        return path
    exif_rot = (im.getexif() or {}).get(0x0112, 1) not in (1, None)
    vis_w = im.width * max(0.01, 1 - crop.get("l", 0) - crop.get("r", 0))
    vis_h = im.height * max(0.01, 1 - crop.get("t", 0) - crop.get("b", 0))
    ratio = max(vis_w / max(1, box_w * 2.0), vis_h / max(1, box_h * 2.0))
    if ratio <= 1.25 and path.suffix.lower() in ok_ext and not exif_rot:
        return path
    if exif_rot:
        im = ImageOps.exif_transpose(im)
    if ratio > 1.25:
        im = im.resize((max(1, round(im.width / ratio)), max(1, round(im.height / ratio))), Image.LANCZOS)
    alpha = im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info)
    media_dir.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha1(f"{path}|{im.size}".encode()).hexdigest()[:16]
    if alpha:
        out = media_dir / f"r-{digest}.png"
        im.convert("RGBA").save(out, optimize=True)
    else:
        out = media_dir / f"r-{digest}.jpg"
        im.convert("RGB").save(out, quality=90, optimize=True)
    return out


def vertical_insets(text, box_h, fx):
    """Top/bottom insets + paragraph spacing so PowerPoint baselines land on Chromium's measured baselines.

    fx.ctx.baseline(p) gives PowerPoint's measured first-baseline offset for the paragraph's style (see calib.py)."""
    paras = text["paragraphs"]
    if "baseline_first" not in paras[0]:
        return text["text_top"], text["pad"]["b"]
    off = [fx.ctx.baseline(p) for p in paras]
    for i, p in enumerate(paras):
        if i == 0:
            p["space_before"] = 0
        else:
            prev = paras[i - 1]
            p["space_before"] = max(0.0, p["baseline_first"] - prev["baseline_last"]
                                    - (prev["line_height"] - off[i - 1]["t"]) - off[i]["t"])
    L0, o = paras[0]["line_height"], off[0]
    block = sum(p["lines"] * p["line_height"] for p in paras) + sum(p["space_before"] for p in paras[1:])
    top = paras[0]["baseline_first"] - o["t"]
    if text["anchor"] == "t":
        return top, text["pad"]["b"]
    if text["anchor"] == "b":
        corr = o["b"] - 2 * L0 - o["t"]
        return text["pad"]["t"], box_h - (top - corr + block)
    corr = o["ctr"] - L0 - o["t"]
    t0 = top - corr
    m = max(0.0, min(t0, box_h - t0 - block))
    return t0 - m, box_h - t0 - block - m


def wrap_delta(wr):
    """Pick a width change inside the range that keeps Chromium's line breaks, favouring the middle."""
    lo, hi = wr["min"], wr["max"]
    d = max(-24.0, min(24.0, (lo + hi) / 2))
    if hi - lo > 1:
        d = max(lo + 0.5, min(hi - 0.5, d))
    return d


def _round_adj(radius, w, h):
    return max(0, min(50000, int(round(radius / max(1e-6, min(w, h)) * 100000))))


# ---------- element builders ----------
def build_text(shapes, el, fx):
    sh = el.get("shape") or {}
    text = el.get("text")
    x, y, w, h = (el["box"][k] for k in ("x", "y", "w", "h"))
    line = sh.get("line")
    bw = line["width"] if line else 0
    fill, shadow, geom = sh.get("fill"), sh.get("shadow"), sh.get("geom", "rect")
    visual = bool(fill or line or shadow or geom != "rect")
    if bw:
        x, y, w, h = x + bw / 2, y + bw / 2, max(1, w - bw), max(1, h - bw)
    a = el.get("opacity", 1)
    insets = [0, 0, 0, 0]
    wrap, anchor = True, "t"
    if text:
        pad = text["pad"]
        anchor = text["anchor"]
        t_in, b_in = vertical_insets(text, el["box"]["h"], fx)
        insets = [pad["l"] - bw / 2, t_in - bw / 2, pad["r"] - bw / 2, b_in - bw / 2]
        if not visual and insets[1] < 0:  # text rises above the box: grow the (invisible) box instead
            y, h, insets[1] = y + insets[1], h - insets[1], 0
        if not visual and insets[3] < 0:
            h, insets[3] = h - insets[3], 0
        aligns = {p["align"] for p in text["paragraphs"]}
        al = next(iter(aligns)) if len(aligns) == 1 else "l"
        wr = text.get("wrap_range")
        if text["single_line"] and "just" not in aligns:
            wrap = False
        elif wr:
            delta = wrap_delta(wr)
            if not visual:
                x -= delta if al == "r" else (delta / 2 if al == "ctr" else 0)
                w += delta
            elif al == "ctr":
                insets[0] -= delta / 2
                insets[2] -= delta / 2
            else:
                insets[0 if al == "r" else 2] -= delta
    if visual:
        prst = {"rect": MSO_SHAPE.RECTANGLE, "roundRect": MSO_SHAPE.ROUNDED_RECTANGLE, "ellipse": MSO_SHAPE.OVAL}[geom]
        shp = shapes.add_shape(prst, Emu(fx.emu(x)), Emu(fx.emu(y)), Emu(fx.emu(w)), Emu(fx.emu(h)))
        _remove_style(shp)
        if geom == "roundRect":
            X.set_geometry(shp._element.spPr, "roundRect", _round_adj(max(0, sh.get("radius", 0) - bw / 2), w, h))
    else:
        shp = shapes.add_textbox(Emu(fx.emu(x)), Emu(fx.emu(y)), Emu(fx.emu(w)), Emu(fx.emu(h)))
    X.set_sppr(shp._element.spPr, X.fill_el(fill, fx.theme, a),
               X.line_el(line, fx.emu, fx.theme, a) if (line or visual) else None,
               X.shadow_el(shadow, fx.emu, w, h, fx.theme, a))
    if el.get("rotation"):
        shp.rotation = el["rotation"]
    txBody = shp._element.find(qn("p:txBody"))
    if text and txBody is not None:
        fx.alpha = a
        X.fill_txbody(txBody, text, fx, wrap, insets, anchor)
        fx.alpha = 1.0
    elif txBody is not None:
        bp = txBody.find(qn("a:bodyPr"))
        if bp is not None:
            bp.set("anchor", "ctr")
    _set_name(shp, el.get("name"))
    return shp


def build_picture(shapes, el, fx):
    ctx = fx.ctx
    im = el["image"]
    x, y, w, h = (el["box"][k] for k in ("x", "y", "w", "h"))
    sh = el.get("shape") or {}
    crop = im.get("crop") or {}
    svg_bytes = None
    if im["kind"] in ("svg", "svg-inline"):
        if not im.get("png"):
            ctx.warnings.append(f"slide {fx.slide_index}: svg without raster fallback skipped ({el.get('src', '')})")
            return None
        src = im["png"]
        svg_bytes = pathlib.Path(im["svg_file"]).read_bytes()
    else:
        if not im.get("path"):
            ctx.warnings.append(f"slide {fx.slide_index}: image skipped, unresolved source ({im.get('src', '')[:80]})")
            return None
        src = prepare_raster(im["path"], w, h, crop, ctx.media_dir)
    pic = shapes.add_picture(str(src), Emu(fx.emu(x)), Emu(fx.emu(y)), Emu(fx.emu(w)), Emu(fx.emu(h)))
    for side in ("l", "t", "r", "b"):
        v = crop.get(side, 0)
        if abs(v) > 1e-4:
            setattr(pic, {"l": "crop_left", "t": "crop_top", "r": "crop_right", "b": "crop_bottom"}[side], v)
    mask = im.get("mask", "rect")
    if mask == "ellipse":
        X.set_geometry(pic._element.spPr, "ellipse")
    elif mask == "roundRect":
        X.set_geometry(pic._element.spPr, "roundRect", _round_adj(im.get("radius", 0), w, h))
    a = el.get("opacity", 1)
    line = sh.get("line")
    X.set_sppr(pic._element.spPr, X.fill_el(sh.get("fill"), fx.theme, a) if sh.get("fill") else None,
               X.line_el(line, fx.emu, fx.theme, a) if line else None,
               X.shadow_el(sh.get("shadow"), fx.emu, w, h, fx.theme, a))
    if svg_bytes:
        X.add_svg_blip(pic, svg_bytes, fx.slide.part, ctx.svg_parts)
    X.set_picture_alpha(pic, a)
    if el.get("rotation"):
        pic.rotation = el["rotation"]
    _set_name(pic, el.get("name"), im.get("alt"))
    return pic


def _cell_border(side_tag, b, fx):
    if not b:
        return None
    ln = X.line_el(b, fx.emu, fx.theme)
    ln.tag = qn(side_tag)
    return ln


def build_table(slide, shapes, el, fx):
    t = el["table"]
    x, y = el["box"]["x"], el["box"]["y"]
    nr, nc = len(t["rows"]), len(t["cols"])
    if not nr or not nc:
        return None
    gf = slide.shapes.add_table(nr, nc, Emu(fx.emu(x)), Emu(fx.emu(y)), Emu(fx.emu(sum(t["cols"]))), Emu(fx.emu(sum(t["rows"]))))
    tbl = gf.table
    for i, cw in enumerate(t["cols"]):
        tbl.columns[i].width = Emu(fx.emu(cw))
    for i, rh in enumerate(t["rows"]):
        tbl.rows[i].height = Emu(fx.emu(rh))
    tblPr = tbl._tbl.tblPr
    for attr in ("firstRow", "bandRow", "firstCol", "lastRow", "lastCol", "bandCol"):
        tblPr.set(attr, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = X.E("a:tableStyleId")
        tblPr.append(sid)
    sid.text = NO_STYLE_TABLE
    for c in t["cells"]:
        r0, c0 = min(c["r"], nr - 1), min(c["c"], nc - 1)
        cell = tbl.cell(r0, c0)
        if c["rs"] > 1 or c["cs"] > 1:
            cell.merge(tbl.cell(min(nr - 1, r0 + c["rs"] - 1), min(nc - 1, c0 + c["cs"] - 1)))
        tc = cell._tc
        text = c.get("text")
        top, bottom = c["pad"]["t"], c["pad"]["b"]
        if text:
            cell_h = sum(t["rows"][r0:r0 + c["rs"]])
            top, bottom = vertical_insets(text, cell_h, fx)
            X.fill_txbody(tc.txBody, text, fx, True, (0, 0, 0, 0), text["anchor"])
        tcPr = tc.get_or_add_tcPr()
        for ch in list(tcPr):
            tcPr.remove(ch)
        for k, v in (("marL", c["pad"]["l"]), ("marR", c["pad"]["r"]), ("marT", top), ("marB", bottom)):
            tcPr.set(k, str(fx.emu(max(0, v))))
        tcPr.set("anchor", c.get("anchor", "t"))
        b = c.get("borders") or {}
        for tag, side in (("a:lnL", "left"), ("a:lnR", "right"), ("a:lnT", "top"), ("a:lnB", "bottom")):
            ln = _cell_border(tag, b.get(side), fx)
            if ln is not None:
                tcPr.append(ln)
        tcPr.append(X.fill_el(c.get("fill"), fx.theme))
    _set_name(gf, el.get("name") or "Table")
    if shapes is not slide.shapes:
        shapes._spTree.append(gf._element)
        shapes._recalculate_extents()
    return gf


def build_chart(shapes, el, fx):
    spec = el["chart"]["spec"]
    font = el["chart"]["font"]
    typ = spec.get("type", "column")
    if typ not in CHART_TYPES:
        fx.ctx.warnings.append(f"slide {fx.slide_index}: unknown chart type {typ!r}")
        return None
    x, y, w, h = (el["box"][k] for k in ("x", "y", "w", "h"))
    if typ == "scatter":
        cd = XyChartData()
        for s in spec["series"]:
            ser = cd.add_series(s["name"])
            for px_, py_ in s["points"]:
                ser.add_data_point(px_, py_)
    else:
        cd = CategoryChartData(number_format=spec.get("number_format", "General"))
        cd.categories = spec["categories"]
        for s in spec["series"]:
            cd.add_series(s["name"], s["values"])
    gf = shapes.add_chart(CHART_TYPES[typ], Emu(fx.emu(x)), Emu(fx.emu(y)), Emu(fx.emu(w)), Emu(fx.emu(h)), cd)
    chart = gf.chart
    typeface, _, _ = fx.font({"family": font["family"], "platform_font": font.get("platform_font"), "weight": 400, "italic": False})
    chart.font.name = typeface
    chart.font.size = Pt(round(fx.pt(spec.get("font_size", font["size"])), 1))
    chart.font.color.rgb = RGBColor.from_string(hex6(spec.get("text_color", font["hex"])))
    chart.has_title = bool(spec.get("title"))
    if spec.get("title"):
        chart.chart_title.text_frame.text = spec["title"]
    palette = [hex6(c) for c in (spec.get("colors") or fx.ctx.palette or ["4472C4", "ED7D31", "A5A5A5", "FFC000", "5B9BD5"])]
    pie = typ in ("pie", "doughnut")
    legend = spec.get("legend", "right" if pie else ("bottom" if len(spec["series"]) > 1 else "none"))
    chart.has_legend = legend != "none"
    if chart.has_legend:
        chart.legend.position = {"bottom": XL_LEGEND_POSITION.BOTTOM, "top": XL_LEGEND_POSITION.TOP, "right": XL_LEGEND_POSITION.RIGHT,
                                 "left": XL_LEGEND_POSITION.LEFT}[legend]
        chart.legend.include_in_layout = False
    plot = chart.plots[0]
    for i, s in enumerate(plot.series):
        col = hex6(spec["series"][i].get("color", palette[i % len(palette)]))
        if pie:
            for j in range(len(spec["categories"])):
                pt_ = s.points[j]
                pt_.format.fill.solid()
                pt_.format.fill.fore_color.rgb = RGBColor.from_string(hex6(palette[j % len(palette)]))
                pt_.format.line.color.rgb = RGBColor.from_string(hex6(spec.get("separator", "FFFFFF")))
        elif typ.startswith("line") or typ == "scatter":
            s.format.line.color.rgb = RGBColor.from_string(col)
            s.format.line.width = Pt(spec.get("line_width", 2.5))
            s.smooth = bool(spec.get("smooth", False))
        else:
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = RGBColor.from_string(col)
    if hasattr(plot, "gap_width") and not pie:
        plot.gap_width = int(spec.get("gap", 60))
        if "stacked" in typ or "percent" in typ:
            plot.overlap = 100
    if typ == "doughnut":
        hole = gf.chart._chartSpace.find(".//" + qn("c:holeSize"))
        if hole is not None:
            hole.set("val", str(int(spec.get("hole", 60))))
    if spec.get("labels"):
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = spec.get("label_format", spec.get("number_format", "General"))
        dl.number_format_is_linked = False
        dl.font.size = Pt(round(fx.pt(spec.get("label_size", font["size"])), 1))
        if pie:
            dl.position = XL_DATA_LABEL_POSITION.OUTSIDE_END if typ == "pie" else XL_DATA_LABEL_POSITION.CENTER
        elif not typ.startswith(("line", "area", "stacked", "percent")):
            dl.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
    if not pie:
        grid = hex6(spec.get("grid_color", "E5E5E5"))
        va, ca = chart.value_axis, chart.category_axis
        va.has_major_gridlines = bool(spec.get("gridlines", True))
        if va.has_major_gridlines:
            va.major_gridlines.format.line.color.rgb = RGBColor.from_string(grid)
            va.major_gridlines.format.line.width = Pt(0.75)
        va.format.line.fill.background()
        va.visible = bool(spec.get("value_axis", True))
        if spec.get("min") is not None:
            va.minimum_scale = spec["min"]
        if spec.get("max") is not None:
            va.maximum_scale = spec["max"]
        va.tick_labels.number_format = spec.get("number_format", "General")
        va.tick_labels.number_format_is_linked = False
        ca.format.line.color.rgb = RGBColor.from_string(grid)
        ca.has_major_gridlines = False
    _set_name(gf, el.get("name") or "Chart")
    return gf


def build(shapes, el, fx, built):
    k = el["kind"]
    shp = None
    if k == "group":
        grp = shapes.add_group_shape()
        for ch in el.get("children", []):
            build(grp.shapes, ch, fx, built)
        if not len(grp.shapes):
            grp._element.getparent().remove(grp._element)
            return None
        _set_name(grp, el.get("name"))
        shp = grp
    elif k in ("text", "shape"):
        shp = build_text(shapes, el, fx)
    elif k == "image":
        shp = build_picture(shapes, el, fx)
    elif k == "table":
        shp = build_table(fx.slide, shapes, el, fx)
    elif k == "chart":
        shp = build_chart(shapes, el, fx)
    if shp is not None:
        built[el["id"]] = (shp, el)
    return shp


def _walk(elements):
    for e in elements:
        yield e
        yield from _walk(e.get("children", []))


def open_template(path):
    """Open a .pptx/.potx as a Presentation (python-pptx rejects the template content type, so patch it in memory)."""
    import io
    import zipfile
    path = pathlib.Path(path)
    if path.suffix.lower() == ".pptx":
        return Presentation(str(path))
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"presentationml.template.main+xml", b"presentationml.presentation.main+xml")
            zout.writestr(item, data)
    buf.seek(0)
    return Presentation(buf)


def reset_core_props(prs, author="", keep_author=False):
    """Stop a template's (or a previous editor's) identity leaking into the deck we save."""
    import datetime
    cp = prs.core_properties
    now = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0, tzinfo=None)
    if not keep_author:
        cp.author = author
        cp.created = now
        cp.revision = 1
    cp.last_modified_by = author
    cp.modified = now


def find_blank_layout(prs):
    best = None
    for layout in prs.slide_layouts:
        n = sum(1 for _ in layout.iter_cloneable_placeholders())
        if layout.name.lower() == "blank" and n == 0:
            return layout
        if best is None or n < best[0]:
            best = (n, layout)
    return best[1]


def _calibrate(ctx, L, enabled):
    from .render import powerpoint_backend, render_powerpoint
    styles = set()
    for s in L["slides"]:
        for el in _walk(s["elements"]):
            texts = [el["text"]] if el.get("text") else []
            if el.get("table"):
                texts += [c["text"] for c in el["table"]["cells"] if c.get("text")]
            for t in texts:
                for p in t["paragraphs"]:
                    if "baseline_first" in p:
                        styles.add(ctx.style_of(p))
    def ppt_render(pptx, out_dir, w, h):
        pathlib.Path(out_dir).mkdir(parents=True, exist_ok=True)
        render_powerpoint(pptx, out_dir, w, h, 240)

    backend = powerpoint_backend()
    renderer = ppt_render if enabled and backend else None
    try:
        ctx.baselines = calib.measure(styles, renderer, backend or "powerpoint")
    except Exception as e:  # noqa: BLE001 - fall back to the model, QA will show any drift
        ctx.warnings.append(f"baseline calibration failed ({e}); using the 0.75 x line-height model")
        ctx.baselines = {}
    return {"styles": len(styles), "measured": sum(1 for v in ctx.baselines.values() if v and v.get("measured"))}


def compile_deck(project, layout_path=None, out=None, brand_path=None, embed=False, calibrate=True, author=""):
    project = pathlib.Path(project).resolve()
    layout_path = pathlib.Path(layout_path or project / "build" / "layout.json")
    L = read_json(layout_path)
    brand_path = pathlib.Path(brand_path or project / "brand" / "brand.json")
    brand = read_json(brand_path) if brand_path.exists() else {}
    template = brand.get("template")
    if template and not pathlib.Path(template).is_absolute():
        template = str((brand_path.parent / template).resolve())
    if template and not pathlib.Path(template).exists():
        log(f"warning: brand template {template} not found; building without it (fix \"template\" in brand.json)")
    if template and pathlib.Path(template).exists():
        prs = open_template(template)
        sldIdLst = prs.slides._sldIdLst
        for sldId in list(sldIdLst):
            prs.part.drop_rel(sldId.rId)
            sldIdLst.remove(sldId)
    else:
        template = None
        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
        slots = brand.get("theme") or {}
        fonts = brand.get("fonts") or {}
        X.patch_theme(prs, slots=slots or None, major=(fonts.get("display") or {}).get("family"),
                      minor=(fonts.get("body") or {}).get("family"), name=brand.get("name"))
    reset_core_props(prs, author)
    k =prs.slide_width / (CANVAS_W * 9525)
    if abs(prs.slide_height / prs.slide_width - 720 / 1280) > 0.01:
        log("warning: template slide size is not 16:9; content is scaled to width")
    media_dir = layout_path.parent / "media"
    ctx = Ctx(prs, k, brand, media_dir, brand_dir=brand_path.parent)
    ctx.palette = [hex6(c) for c in (brand.get("chart_palette") or [v for s, v in sorted((brand.get("theme") or {}).items()) if s.startswith("accent")])]
    blank = find_blank_layout(prs)
    report = {"slides": [], "fonts": {}, "warnings": ctx.warnings, "calibration": _calibrate(ctx, L, calibrate)}

    for s in L["slides"]:
        slide = prs.slides.add_slide(blank)
        for ph in list(slide.placeholders):
            ph._element.getparent().remove(ph._element)
        if template:
            slide._element.set("showMasterSp", "0")
        fx = Fx(ctx, slide)
        fx.slide_index = s["index"]
        bg = s.get("background") or {}
        if bg.get("image") and (bg["image"]["image"].get("path") or bg["image"]["image"].get("png")):
            im = bg["image"]["image"]
            src = im.get("png") or prepare_raster(im["path"], 1280, 720, im.get("crop") or {}, media_dir)
            _, rid = slide.part.get_or_add_image_part(str(src))
            X.set_background(slide, blip_rid=rid, crop=im.get("crop"))
            if bg.get("fill") and bg["fill"]["type"] == "gradient":
                build_text(slide.shapes, {"id": "bg-overlay", "kind": "shape", "name": "Background overlay",
                                          "box": {"x": 0, "y": 0, "w": 1280, "h": 720}, "opacity": 1,
                                          "shape": {"fill": bg["fill"], "geom": "rect"}}, fx)
        elif bg.get("fill"):
            X.set_background(slide, fill=bg["fill"], theme=fx.theme)
        built = {}
        for el in s["elements"]:
            build(slide.shapes, el, fx, built)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
        trans = s.get("transition")
        if trans:
            try:
                X.set_transition(slide, trans)
            except ValueError as e:
                ctx.warnings.append(f"slide {s['index']}: {e}")
        anims = []
        for i, el in enumerate(_walk(s["elements"])):
            if el.get("anim") and el["id"] in built:
                shp, _ = built[el["id"]]
                a = el["anim"]
                anims.append((a.get("order") if a.get("order") is not None else 10_000, i, {
                    "spid": shp.shape_id, "is_sp": shp._element.tag == qn("p:sp"), "effect": a["effect"],
                    "trigger": a.get("trigger", "click"), "delay": a.get("delay", 0), "duration": a.get("duration")}))
        anims.sort(key=lambda t: (t[0], t[1]))
        if anims:
            try:
                X.set_timing(slide, [a for _, _, a in anims])
            except ValueError as e:
                ctx.warnings.append(f"slide {s['index']}: {e}")
        report["slides"].append({"index": s["index"], "shapes": len(slide.shapes), "animations": len(anims),
                                 "transition": (trans or {}).get("type")})

    sysdir = str(system_font_dir() or "").lower()
    for (typeface, bold, italic), face in ctx.used_faces.items():
        entry = report["fonts"].setdefault(typeface, {"installed": True, "system": face["path"].lower().startswith(sysdir) if sysdir else False,
                                                      "files": {}, "embeddable": embeddable(face)})
        style = ("boldItalic" if italic else "bold") if bold else ("italic" if italic else "regular")
        entry["files"][style] = face["path"]
    for fam, n in ctx.missing.items():
        report["fonts"][fam] = {"installed": False, "runs": n}
    if embed:
        targets = {tf: info["files"] for tf, info in report["fonts"].items()
                   if info.get("installed") and not info.get("system") and info.get("embeddable")}
        report["embedded_fonts"] = X.embed_fonts(prs, targets)
    out = pathlib.Path(out or project / "deck.pptx")
    prs.save(out)
    report["output"] = str(out)
    # lets `inspect`/Edit mode tell whether someone changed the file in PowerPoint after this build
    report["sha256"] = hashlib.sha256(out.read_bytes()).hexdigest()
    write_json(layout_path.parent / "compile.json", report)
    return report
