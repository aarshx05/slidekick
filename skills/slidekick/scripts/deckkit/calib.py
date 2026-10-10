"""Measure where PowerPoint places the first baseline for each text style, instead of guessing.

PowerPoint's exact line spacing puts the baseline at 0.75 x spacing when spacing >= ~1.3 x font size, but switches
between regimes non-smoothly for tighter headings. So for every (typeface, bold, italic, size, spacing) the deck uses,
we render 'HHHH' probes once through real PowerPoint and cache the measured offsets (top / centre / bottom anchors).
"""
import json
import pathlib
import tempfile

import numpy as np
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

from .util import cache_dir, log

SCALE = 3
VERSION = 1


def cache_path(backend="powerpoint"):
    """Baselines differ between PowerPoint for Windows and for Mac, so each keeps its own cache."""
    return cache_dir() / ("ppt-baselines.json" if backend == "powerpoint" else f"ppt-baselines-{backend}.json")


def key(typeface, bold, italic, sz, spc):
    return f"v{VERSION}|{typeface}|{int(bold)}|{int(italic)}|{int(sz)}|{int(spc)}"


def _load(path):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:  # noqa: BLE001
        return {}


def model(spc_hpt):
    """Fallback when PowerPoint is unavailable: first baseline at 0.75 x line spacing (px)."""
    L = spc_hpt / 100 / 0.75
    return {"t": 0.75 * L, "ctr": 0.75 * L + L, "b": 0.75 * L + 2 * L, "measured": False}


def _probe_deck(styles, path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(1280 * 9525), Emu(720 * 9525)
    placed, slide, x, y, col_w = [], None, 20, 20, 0
    for st in styles:
        typeface, bold, italic, sz, spc = st
        size_px, L = sz / 100 / 0.75, spc / 100 / 0.75
        h, gap, w = 3 * L, max(size_px, L) + 24, max(160, size_px * 3.6)
        for anchor in ("t", "ctr", "b"):
            if slide is None or y + h + gap > 720:
                if slide is None or x + col_w + w + 40 > 1280:
                    slide, x, col_w = prs.slides.add_slide(prs.slide_layouts[6]), 20, 0
                else:
                    x += col_w + 40
                    col_w = 0
                y = max(20, size_px)
            tb = slide.shapes.add_textbox(Emu(int(x * 9525)), Emu(int(y * 9525)), Emu(int(w * 9525)), Emu(int(h * 9525)))
            body = tb.text_frame._txBody
            bp = body.find(qn("a:bodyPr"))
            for ch in list(bp):
                bp.remove(ch)
            for k in ("lIns", "tIns", "rIns", "bIns"):
                bp.set(k, "0")
            bp.set("wrap", "none")
            bp.set("anchor", anchor)
            etree.SubElement(bp, qn("a:noAutofit"))
            p = tb.text_frame.paragraphs[0]._p
            ppr = etree.SubElement(p, qn("a:pPr"))
            ls = etree.SubElement(etree.SubElement(ppr, qn("a:lnSpc")), qn("a:spcPts"))
            ls.set("val", str(int(spc)))
            r = etree.SubElement(p, qn("a:r"))
            rpr = etree.SubElement(r, qn("a:rPr"), lang="en-US", sz=str(int(sz)), b=str(int(bold)), i=str(int(italic)), kern="0")
            fill = etree.SubElement(etree.SubElement(rpr, qn("a:solidFill")), qn("a:srgbClr"))
            fill.set("val", "000000")
            etree.SubElement(rpr, qn("a:latin"), typeface=typeface)
            etree.SubElement(r, qn("a:t")).text = "HHHH"
            placed.append({"style": st, "anchor": anchor, "slide": len(prs.slides), "x": x, "y": y, "w": w, "h": h, "size": size_px})
            y += h + gap
            col_w = max(col_w, w)
    prs.save(path)
    return placed


def measure(styles, renderer, backend="powerpoint"):
    """styles: iterable of (typeface, bold, italic, sz, spc) tuples. Returns {key: offsets} (cached per backend)."""
    path = cache_path(backend)
    cache = _load(path)
    want = sorted({tuple(s) for s in styles})
    missing = [s for s in want if key(*s) not in cache]
    if missing and renderer is not None:
        log(f"calibrating {len(missing)} text style(s) against PowerPoint")
        with tempfile.TemporaryDirectory() as tmp:
            tmp = pathlib.Path(tmp)
            placed = _probe_deck(missing, tmp / "probe.pptx")
            renderer(tmp / "probe.pptx", tmp / "render", 1280 * SCALE, 720 * SCALE)
            imgs = {}
            for pr in placed:
                im = imgs.get(pr["slide"])
                if im is None:
                    im = imgs[pr["slide"]] = np.asarray(Image.open(tmp / "render" / f"slide-{pr['slide']:02d}.png").convert("L"))
                y0 = int((pr["y"] - pr["size"]) * SCALE)
                y1 = int((pr["y"] + pr["h"] + pr["size"]) * SCALE)
                x0, x1 = int(pr["x"] * SCALE), int((pr["x"] + pr["w"]) * SCALE)
                ink = np.where((im[max(0, y0):y1, x0:x1] < 128).sum(axis=1) > 2)[0]
                k = key(*pr["style"])
                entry = cache.setdefault(k, {"measured": True})
                if len(ink):
                    entry[pr["anchor"]] = round((max(0, y0) + ink[-1] + 1) / SCALE - pr["y"], 3)
            for s in missing:
                e = cache.get(key(*s), {})
                if not all(a in e for a in ("t", "ctr", "b")):
                    cache[key(*s)] = model(s[4])
        try:
            path.write_text(json.dumps(cache, indent=0), encoding="utf-8")
        except OSError:
            pass
    return {key(*s): cache.get(key(*s)) or model(s[4]) for s in want}
