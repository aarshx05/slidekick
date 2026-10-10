"""In-place editing of an existing .pptx: apply edits.json (a list of operations) without regenerating the deck.

Everything not named by an operation is left byte-for-byte alone: masters, layouts, other slides, untouched shapes and
the formatting of edited text. Slide references are the slide numbers shown by `deck.py inspect` (the original deck),
or labels given to slides created by `import`/`duplicate`. Shape references are an id (int) or an exact name.
All operations succeed or nothing is written."""
import copy
import pathlib
import re
import shutil

from pptx.chart.data import CategoryChartData, XyChartData
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import PartFactory
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.util import Emu

from .util import CANVAS_W, hex6, read_json

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
P14 = "{http://schemas.microsoft.com/office/powerpoint/2010/main}"
SLOTS = {"dk1", "lt1", "dk2", "lt2", "tx1", "bg1", "tx2", "bg2", "accent1", "accent2", "accent3", "accent4", "accent5",
         "accent6", "hlink", "folHlink"}
SHAREABLE = ("image/", "video/", "audio/")


class EditError(Exception):
    pass


# ---------- addressing ----------
class Deck:
    def __init__(self, prs):
        self.prs = prs
        self.refs = {i: s for i, s in enumerate(prs.slides, 1)}  # original numbering, stable across ops
        self.deleted = set()
        self.k = CANVAS_W * 9525 / prs.slide_width  # canvas px -> EMU factor is 9525 / k

    def slide(self, ref, op):
        if isinstance(ref, str) and ref.isdigit():
            ref = int(ref)
        s = self.refs.get(ref)
        if s is None:
            raise EditError(f"{op}: unknown slide {ref!r} (use the numbers from `inspect`, or a label set with \"as\")")
        if id(s) in self.deleted:
            raise EditError(f"{op}: slide {ref!r} was deleted by an earlier operation")
        return s

    def emu(self, px):
        return Emu(int(round(px * 9525 / self.k)))

    def position(self, slide):
        for i, s in enumerate(self.prs.slides):
            if s.slide_id == slide.slide_id:
                return i
        raise EditError("slide not in deck")

    def place(self, slide, op, after=None, before=None, at=None):
        """Move `slide` to a position: after/before a slide ref, or at "start"/"end"/1-based index."""
        lst = self.prs.slides._sldIdLst
        el = next(x for x in lst if int(x.get("id")) == slide.slide_id)
        if after is None and before is None and at is None:
            return
        lst.remove(el)
        if after is not None or before is not None:
            anchor = self.slide(after if after is not None else before, op)
            ael = next(x for x in lst if int(x.get("id")) == anchor.slide_id)
            idx = list(lst).index(ael) + (1 if after is not None else 0)
        elif at == "start":
            idx = 0
        elif at == "end":
            idx = len(lst)
        else:
            idx = max(0, min(len(lst), int(at) - 1))
        lst.insert(idx, el)


def iter_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh._element.tag == P + "grpSp":
            yield from iter_shapes(sh.shapes)


def find_shape(slide, ref, op):
    if ref == "@title":
        t = slide.shapes.title
        if t is None:
            raise EditError(f"{op}: slide has no title placeholder")
        return t
    if isinstance(ref, str) and ref.isdigit():
        ref = int(ref)
    hits = [s for s in iter_shapes(slide.shapes) if (s.shape_id == ref if isinstance(ref, int) else s.name == ref)]
    if not hits:
        names = ", ".join(f"#{s.shape_id} {s.name!r}" for s in iter_shapes(slide.shapes))
        raise EditError(f"{op}: no shape {ref!r} on that slide. Shapes: {names}")
    if len(hits) > 1:
        raise EditError(f"{op}: {len(hits)} shapes are named {ref!r}; use the id instead: "
                        + ", ".join(f"#{s.shape_id}" for s in hits))
    return hits[0]


# ---------- text ----------
def _runs_of(p):
    return [r for r in p if r.tag in (A + "r", A + "fld")]


def _new_run(rPr, text):
    r = copy.deepcopy(rPr) if rPr is not None else None
    run = etree_el("a:r")
    if r is not None:
        r.tag = A + "rPr"
        run.append(r)
    t = etree_el("a:t")
    t.text = text
    run.append(t)
    return run


def etree_el(tag):
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls
    pfx, local = tag.split(":")
    return parse_xml(f"<{tag} {nsdecls(pfx)}/>")


def _apply_run_style(rPr_parent_run, style):
    """Apply {bold, italic, color, size_pt, font, underline} to an a:r (creating rPr as needed)."""
    rPr = rPr_parent_run.find(A + "rPr")
    if rPr is None:
        rPr = etree_el("a:rPr")
        rPr_parent_run.insert(0, rPr)
    _style_rpr(rPr, style)


def _color_el(value):
    v = str(value).strip()
    if v.lstrip("#").lower() in ("none", "transparent"):
        return etree_el("a:noFill")
    sf = etree_el("a:solidFill")
    if v in SLOTS:
        c = etree_el("a:schemeClr")
        c.set("val", v)
    else:
        c = etree_el("a:srgbClr")
        c.set("val", hex6(v))
    sf.append(c)
    return sf


def _style_rpr(rPr, style):
    if "size_pt" in style:
        rPr.set("sz", str(int(round(float(style["size_pt"]) * 100))))
    if "bold" in style:
        rPr.set("b", "1" if style["bold"] else "0")
    if "italic" in style:
        rPr.set("i", "1" if style["italic"] else "0")
    if "underline" in style:
        rPr.set("u", "sng" if style["underline"] else "none")
    if "color" in style:
        for old in rPr.findall(A + "solidFill") + rPr.findall(A + "noFill") + rPr.findall(A + "gradFill"):
            rPr.remove(old)
        fill = _color_el(style["color"])
        # fill goes after a:ln, before effect/latin children
        ln = rPr.find(A + "ln")
        rPr.insert(list(rPr).index(ln) + 1 if ln is not None else 0, fill)
    if "font" in style:
        old = rPr.find(A + "latin")
        if old is not None:
            rPr.remove(old)
        lat = etree_el("a:latin")
        lat.set("typeface", style["font"])
        # schema order: … latin, ea, cs, sym, hlinkClick, hlinkMouseOver, rtl, extLst
        nxt = next((c for c in rPr if c.tag.split("}")[-1] in ("ea", "cs", "sym", "hlinkClick", "hlinkMouseOver", "rtl", "extLst")), None)
        (nxt.addprevious(lat) if nxt is not None else rPr.append(lat))


def _fill_paragraph(p, item):
    """Replace the runs of a:p with `item` (str or {"text"|"runs", "level"}), keeping the first run's formatting."""
    runs = _runs_of(p)
    base = None
    if runs:
        base = runs[0].find(A + "rPr")
    elif p.find(A + "endParaRPr") is not None:
        base = p.find(A + "endParaRPr")
    for r in list(p):
        if r.tag in (A + "r", A + "fld", A + "br"):
            p.remove(r)
    end = p.find(A + "endParaRPr")
    spec = item if isinstance(item, dict) else {"text": str(item)}
    if "level" in spec:
        pPr = p.find(A + "pPr")
        if pPr is None:
            pPr = etree_el("a:pPr")
            p.insert(0, pPr)
        pPr.set("lvl", str(int(spec["level"])))
    pieces = spec.get("runs") or [{"text": spec.get("text", "")}]
    for piece in pieces:
        if isinstance(piece, str):
            piece = {"text": piece}
        lines = str(piece.get("text", "")).split("\v")
        for j, line in enumerate(lines):
            if j:
                br = etree_el("a:br")
                if base is not None:
                    b = copy.deepcopy(base)
                    b.tag = A + "rPr"
                    br.append(b)
                (end.addprevious(br) if end is not None else p.append(br))
            run = _new_run(base, line)
            style = {k: v for k, v in piece.items() if k != "text"}
            if style:
                _apply_run_style(run, style)
            (end.addprevious(run) if end is not None else p.append(run))


def set_text(txBody, value):
    """Set a text body's paragraphs. value: str ("\\n" = new paragraph, "\\v" = line break) or a list of paragraphs."""
    items = value.split("\n") if isinstance(value, str) else list(value)
    if not items:
        items = [""]
    paras = txBody.findall(A + "p")
    if not paras:
        paras = [etree_el("a:p")]
        txBody.append(paras[0])
    template = copy.deepcopy(paras[-1])  # extra paragraphs take the last paragraph's formatting
    for i, item in enumerate(items):
        if i < len(paras):
            p = paras[i]
        else:
            p = copy.deepcopy(template)
            txBody.findall(A + "p")[-1].addnext(p)
        _fill_paragraph(p, item)
    for p in txBody.findall(A + "p")[len(items):]:
        txBody.remove(p)


def replace_in_paragraph(p, pattern, repl, regex):
    """Find/replace inside one a:p across run boundaries, keeping each run's formatting. Returns the count."""
    count = 0
    pos = 0
    while True:
        runs = [r for r in p if r.tag == A + "r"]
        texts = [(r.find(A + "t").text or "") if r.find(A + "t") is not None else "" for r in runs]
        full = "".join(texts)
        m = pattern.search(full, pos)
        if not m or m.end() == m.start():
            return count
        new = m.expand(repl) if regex else repl
        s, e = m.start(), m.end()
        offs, o = [], 0
        for t in texts:
            offs.append(o)
            o += len(t)
        i = max(j for j in range(len(runs)) if offs[j] <= s)
        j = max(jj for jj in range(len(runs)) if offs[jj] < e)
        ti = runs[i].find(A + "t")
        tj = runs[j].find(A + "t")
        tail = texts[j][e - offs[j]:]
        ti.text = texts[i][:s - offs[i]] + new + (tail if i == j else "")
        if i != j:
            for k in range(i + 1, j):
                p.remove(runs[k])
            if tail:
                tj.text = tail
            else:
                p.remove(runs[j])
        count += 1
        pos = s + len(new)


def _text_bodies(slide, notes=False):
    for sh in iter_shapes(slide.shapes):
        if sh.has_text_frame:
            yield sh.text_frame._txBody
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    yield c.text_frame._txBody
        if getattr(sh, "has_chart", False) and sh.has_chart and sh.chart.has_title:
            yield sh.chart.chart_title.text_frame._txBody
    if notes and slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
        yield slide.notes_slide.notes_text_frame._txBody


# ---------- parts: cross-package slide cloning ----------
def _next_partname(pkg, src_name, reserved):
    existing = {str(p.partname) for p in pkg.iter_parts()} | reserved
    m = re.match(r"^(.*?)(\d*)(\.[^./]+)$", str(src_name))
    base, ext = m.group(1), m.group(3)
    i = 1
    while f"{base}{i}{ext}" in existing:
        i += 1
    name = f"{base}{i}{ext}"
    reserved.add(name)
    return PackURI(name)


def _remap_rids(root, rid_map, drop):
    for el in list(root.iter()):
        for k, v in list(el.attrib.items()):
            if not k.startswith("{%s}" % R_NS):
                continue
            if v in drop:
                if el.tag in (A + "hlinkClick", A + "hlinkHover"):
                    el.getparent().remove(el)
                    break
                del el.attrib[k]
            elif v in rid_map:
                el.set(k, rid_map[v])


def _clone_part(src, dst_pkg, memo):
    key = id(src)
    if key in memo["parts"]:
        return memo["parts"][key]
    if src.package is dst_pkg and src.content_type.startswith(SHAREABLE):
        memo["parts"][key] = src
        return src
    new = PartFactory(_next_partname(dst_pkg, src.partname, memo["reserved"]), src.content_type, dst_pkg, src.blob)
    memo["parts"][key] = new
    rid_map, drop = {}, set()
    for rid, rel in list(src.rels.items()):
        if rel.is_external:
            rid_map[rid] = new.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif rel.reltype in (RT.SLIDE, RT.NOTES_SLIDE, RT.SLIDE_LAYOUT, RT.NOTES_MASTER):
            drop.add(rid)
        else:
            rid_map[rid] = new.relate_to(_clone_part(rel.target_part, dst_pkg, memo), rel.reltype)
    if hasattr(new, "_element"):
        _remap_rids(new._element, rid_map, drop)
    return new


def _layout_for(prs, src_slide, mode):
    from .compile import find_blank_layout
    if mode == "match":
        name = src_slide.slide_layout.name
        for layout in prs.slide_layouts:
            if layout.name == name:
                return layout
    return find_blank_layout(prs)


def clone_slide(prs, src_slide, layout_mode="blank", memo=None):
    """Copy a slide (from this or another presentation) to the end of `prs`, with its pictures, charts, links,
    background, transition, animations and notes. Returns the new slide."""
    memo = memo if memo is not None else {"parts": {}, "reserved": set()}
    layout = _layout_for(prs, src_slide, layout_mode)
    new = prs.slides.add_slide(layout)
    nel, sel = new._element, src_slide._element
    for child in list(nel):
        nel.remove(child)
    for child in sel:
        nel.append(copy.deepcopy(child))
    for k, v in sel.attrib.items():
        nel.set(k, v)
    for cached in ("shapes", "placeholders"):  # python-pptx lazily caches these against the replaced cSld
        new.__dict__.pop(cached, None)
    rid_map, drop = {}, set()
    for rid, rel in list(src_slide.part.rels.items()):
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
            continue
        if rel.is_external:
            rid_map[rid] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif rel.reltype == RT.SLIDE:
            drop.add(rid)  # "go to slide" links can't follow a slide into another deck
        else:
            rid_map[rid] = new.part.relate_to(_clone_part(rel.target_part, prs.part.package, memo), rel.reltype)
    _remap_rids(nel, rid_map, drop)
    if src_slide.has_notes_slide:
        tf = src_slide.notes_slide.notes_text_frame
        if tf is not None and tf.text.strip():
            set_notes(prs, new, tf.text, donor=src_slide)
    return new


NOTES_BODY = ('<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
              'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:nvSpPr><p:cNvPr id="3" name="Notes Placeholder 2"/>'
              '<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph type="body" idx="1"/></p:nvPr></p:nvSpPr>'
              '<p:spPr><a:xfrm><a:off x="685800" y="4400550"/><a:ext cx="5486400" cy="3600450"/></a:xfrm></p:spPr>'
              '<p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>')


def set_notes(prs, slide, text, donor=None):
    """Speaker notes that work even when the notes master has no body placeholder (common in exported decks)."""
    ns = slide.notes_slide
    if ns.notes_text_frame is None:
        from pptx.oxml import parse_xml
        tmpl = None
        for cand in ([donor] if donor is not None else []) + list(prs.slides):
            if cand is not None and cand.has_notes_slide and cand.notes_slide.notes_text_frame is not None:
                tmpl = copy.deepcopy(cand.notes_slide.notes_placeholder._element)
                break
        if tmpl is None:
            tmpl = parse_xml(NOTES_BODY)
        ns.shapes._spTree.append(tmpl)
    ns.notes_text_frame.text = text


def delete_slide(deck, slide):
    prs = deck.prs
    lst = prs.slides._sldIdLst
    for sid in list(lst):
        if prs.part.related_part(sid.rId) is slide.part:
            rid = sid.rId
            lst.remove(sid)
            # custom shows reference slides by rId of the presentation part
            for el in prs.part._element.iter(P + "sld"):
                if el.get(qn("r:id")) == rid:
                    el.getparent().remove(el)
            prs.part.drop_rel(rid)
            break
    # links from other slides to this one would keep it alive as an orphan part
    for other in prs.slides:
        for rid, rel in list(other.part.rels.items()):
            if not rel.is_external and rel.reltype == RT.SLIDE and rel.target_part is slide.part:
                _remap_rids(other._element, {}, {rid})
                other.part.drop_rel(rid)
    deck.deleted.add(id(slide))


def sync_sections(prs, old_membership):
    """Keep PowerPoint sections valid after slides were added, removed or moved (sections must stay contiguous)."""
    secs = list(prs.part._element.iter(P14 + "section"))
    if not secs:
        return
    import bisect
    order = [s.slide_id for s in prs.slides]
    want = [old_membership.get(sid) for sid in order]
    # keep the longest run of slides whose section order is still non-decreasing; moved and new slides
    # join the section of the slide before them (or after them, at the very start)
    known = [i for i, w in enumerate(want) if w is not None]
    tails, tails_idx, prev = [], [], {}
    for i in known:
        k = bisect.bisect_right(tails, want[i])
        prev[i] = tails_idx[k - 1] if k else None
        if k == len(tails):
            tails.append(want[i])
            tails_idx.append(i)
        else:
            tails[k], tails_idx[k] = want[i], i
    keep, i = set(), (tails_idx[-1] if tails_idx else None)
    while i is not None:
        keep.add(i)
        i = prev[i]
    sec_of = [want[i] if i in keep else None for i in range(len(order))]
    first = next((v for v in sec_of if v is not None), 0)
    cur = first
    assign = {i: [] for i in range(len(secs))}
    for i, sid in enumerate(order):
        if sec_of[i] is not None:
            cur = sec_of[i]
        assign[cur].append(sid)
    from lxml import etree
    for i, sec in enumerate(secs):
        lst = sec.find(P14 + "sldIdLst")
        if lst is None:
            lst = etree.SubElement(sec, P14 + "sldIdLst")
        for x in list(lst):
            lst.remove(x)
        for sid in assign[i]:
            etree.SubElement(lst, P14 + "sldId").set("id", str(sid))


def _section_membership(prs):
    out = {}
    for i, sec in enumerate(prs.part._element.iter(P14 + "section")):
        for x in sec.iter(P14 + "sldId"):
            out[int(x.get("id"))] = i
    return out


# ---------- shapes ----------
def _drop_shape_animations(slide, spid, warnings):
    timing = slide._element.find(P + "timing")
    if timing is None:
        return
    hits = [t for t in timing.iter(P + "spTgt") if t.get("spid") == str(spid)]
    if not hits:
        return
    effects = []
    for t in hits:
        node = t
        # climb to the effect's own p:par (the one whose cTn has a presetClass)
        while node is not None and not (node.tag == P + "par" and node.find(P + "cTn") is not None
                                        and node.find(P + "cTn").get("presetClass")):
            node = node.getparent()
        if node is None:
            timing.getparent().remove(timing)
            warnings.append("removed all animations on a slide (could not isolate the deleted shape's effect)")
            return
        if not any(node is e for e in effects):  # one effect can target the shape several times
            effects.append(node)
    for node in effects:
        parent = node.getparent()
        parent.remove(node)
        # prune now-empty click groups
        while parent is not None and parent.tag == P + "childTnLst" and not len(parent):
            par = parent.getparent().getparent()  # childTnLst -> cTn -> par
            if par is None or par.tag != P + "par":
                break
            gp = par.getparent()
            gp.remove(par)
            parent = gp
    for b in list(timing.iter(P + "bldP")):
        if b.get("spid") == str(spid):
            b.getparent().remove(b)
    if not list(timing.iter(P + "spTgt")):
        timing.getparent().remove(timing)


def _blip_fill(el):
    for path in (f"{P}blipFill", f"{P}spPr/{A}blipFill"):
        bf = el.find(path)
        if bf is not None:
            return bf
    return None


def replace_image(slide, shape, path, fit, focus, warnings):
    from PIL import Image
    path = pathlib.Path(path)
    if not path.exists():
        raise EditError(f"image: {path} not found")
    if path.suffix.lower() == ".svg":
        raise EditError("image: SVG replacement isn't supported in place; export it to PNG first (2x its display size)")
    bf = _blip_fill(shape._element)
    if bf is None:
        raise EditError(f"image: shape #{shape.shape_id} {shape.name!r} has no picture to replace")
    _, rid = slide.part.get_or_add_image_part(str(path))
    blip = bf.find(A + "blip")
    blip.set(qn("r:embed"), rid)
    ext = blip.find(A + "extLst")
    if ext is not None:  # drop an old SVG twin, it would win over the new raster
        for e in list(ext):
            if "svgBlip" in "".join(c.tag for c in e):
                ext.remove(e)
    with Image.open(path) as im:
        iw, ih = im.size
    old = bf.find(A + "srcRect")
    if old is not None:
        bf.remove(old)
    bw, bh = shape.width, shape.height
    if fit == "contain":
        s = min(bw / iw, bh / ih)
        nw, nh = int(iw * s), int(ih * s)
        shape.left, shape.top = shape.left + (bw - nw) // 2, shape.top + (bh - nh) // 2
        shape.width, shape.height = nw, nh
    elif fit == "cover":
        ai, ab = iw / ih, bw / bh
        fx, fy = focus
        src = etree_el("a:srcRect")
        if ai > ab:
            c = 1 - ab / ai
            l_ = c * fx
            src.set("l", str(int(round(l_ * 100000))))
            src.set("r", str(int(round((c - l_) * 100000))))
        elif ai < ab:
            c = 1 - ai / ab
            t_ = c * fy
            src.set("t", str(int(round(t_ * 100000))))
            src.set("b", str(int(round((c - t_) * 100000))))
        blip.addnext(src)
    elif fit != "stretch":
        raise EditError(f"image: fit must be cover|contain|stretch, not {fit!r}")
    if (iw / max(1, shape.width / 9525)) < 1.0:
        warnings.append(f"image {path.name} is smaller than its frame — it will look soft")


# ---------- operations ----------
def op_text(d, o, w):
    slide = d.slide(o["slide"], "text")
    sh = find_shape(slide, o["shape"], "text")
    if not sh.has_text_frame:
        raise EditError(f"text: #{sh.shape_id} {sh.name!r} has no text frame")
    set_text(sh.text_frame._txBody, o["text"])
    return f"text      slide {o['slide']} #{sh.shape_id} {sh.name!r}"


def op_replace(d, o, w):
    flags = 0 if o.get("case", True) else re.I
    regex = bool(o.get("regex"))
    pattern = re.compile(o["find"] if regex else re.escape(o["find"]), flags)
    slides = [d.slide(o["slide"], "replace")] if o.get("slide") is not None else \
        [s for s in d.prs.slides if id(s) not in d.deleted]
    n = 0
    for s in slides:
        for body in _text_bodies(s, notes=o.get("notes", False)):
            for p in body.findall(A + "p"):
                n += replace_in_paragraph(p, pattern, o["with"], regex)
    if not n and not o.get("optional"):
        raise EditError(f"replace: {o['find']!r} not found (set \"optional\": true to allow no match)")
    return f"replace   {o['find']!r} -> {o['with']!r}: {n} occurrence(s)"


def op_notes(d, o, w):
    slide = d.slide(o["slide"], "notes")
    set_notes(d.prs, slide, o["text"])
    return f"notes     slide {o['slide']}"


def op_style(d, o, w):
    slide = d.slide(o["slide"], "style")
    sh = find_shape(slide, o["shape"], "style")
    run_keys = {"size_pt", "bold", "italic", "underline", "color", "font"}
    style = {k: v for k, v in o.items() if k in run_keys}
    if style:
        if not sh.has_text_frame:
            raise EditError(f"style: #{sh.shape_id} has no text")
        paras = sh.text_frame._txBody.findall(A + "p")
        pick = o.get("paragraphs")
        for i, p in enumerate(paras):
            if pick is not None and i not in pick:
                continue
            for r in _runs_of(p):
                _apply_run_style(r, style)
            end = p.find(A + "endParaRPr")
            if end is not None:
                _style_rpr(end, style)
    if "align" in o:
        al = {"left": "l", "center": "ctr", "centre": "ctr", "right": "r", "justify": "just"}[o["align"]]
        for p in sh.text_frame._txBody.findall(A + "p"):
            pPr = p.find(A + "pPr")
            if pPr is None:
                pPr = etree_el("a:pPr")
                p.insert(0, pPr)
            pPr.set("algn", al)
    if "autofit" in o:
        bp = sh.text_frame._txBody.find(A + "bodyPr")
        for t in ("noAutofit", "normAutofit", "spAutoFit"):
            x = bp.find(A + t)
            if x is not None:
                bp.remove(x)
        tag = {"none": "a:noAutofit", "shrink": "a:normAutofit", "resize": "a:spAutoFit"}[o["autofit"]]
        warp = bp.find(A + "prstTxWarp")
        (warp.addnext(etree_el(tag)) if warp is not None else bp.insert(0, etree_el(tag)))
    if "fill" in o or "line" in o:
        spPr = sh._element.find(P + "spPr")
        if spPr is None:
            if sh._element.tag == P + "grpSp":
                kids = ", ".join(f"#{c.shape_id} {c.name!r}" for c in sh.shapes)
                raise EditError(f"style: #{sh.shape_id} {sh.name!r} is a group; style one of its members instead: {kids}")
            raise EditError(f"style: #{sh.shape_id} can't take a fill/line")
        if "fill" in o:
            for t in ("noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill"):
                x = spPr.find(A + t)
                if x is not None:
                    spPr.remove(x)
            geom = spPr.find(A + "prstGeom")
            if geom is None:
                geom = spPr.find(A + "custGeom")
            fill = _color_el(o["fill"])
            (geom.addnext(fill) if geom is not None else spPr.append(fill))
        if "line" in o:
            ln = spPr.find(A + "ln")
            if ln is None:
                ln = etree_el("a:ln")
                spPr.append(ln)
            for t in ("noFill", "solidFill", "gradFill"):
                x = ln.find(A + t)
                if x is not None:
                    ln.remove(x)
            ln.insert(0, _color_el(o["line"]))
    return f"style     slide {o['slide']} #{sh.shape_id} {sh.name!r}: " + ", ".join(k for k in o if k not in ("op", "slide", "shape"))


def op_geometry(d, o, w):
    slide = d.slide(o["slide"], "geometry")
    sh = find_shape(slide, o["shape"], "geometry")
    if sh._element.getparent().tag == P + "grpSp":
        raise EditError("geometry: shape is inside a group; move the group instead")
    for key, attr in (("x", "left"), ("y", "top"), ("w", "width"), ("h", "height")):
        if key in o:
            setattr(sh, attr, d.emu(o[key]))
    return f"geometry  slide {o['slide']} #{sh.shape_id} {sh.name!r}"


def op_delete_shape(d, o, w):
    slide = d.slide(o["slide"], "delete_shape")
    sh = find_shape(slide, o["shape"], "delete_shape")
    sid = sh.shape_id
    sh._element.getparent().remove(sh._element)
    _drop_shape_animations(slide, sid, w)
    return f"delete    slide {o['slide']} shape #{sid}"


def op_image(d, o, w):
    slide = d.slide(o["slide"], "image")
    sh = find_shape(slide, o["shape"], "image")
    replace_image(slide, sh, o["src"], o.get("fit", "cover"), tuple(o.get("focus", (0.5, 0.5))), w)
    if o.get("alt") is not None:
        sh._element.find(f".//{P}cNvPr").set("descr", o["alt"])
    return f"image     slide {o['slide']} #{sh.shape_id} <- {pathlib.Path(o['src']).name} ({o.get('fit', 'cover')})"


def op_chart(d, o, w):
    slide = d.slide(o["slide"], "chart")
    sh = find_shape(slide, o["shape"], "chart")
    if not getattr(sh, "has_chart", False) or not sh.has_chart:
        raise EditError(f"chart: #{sh.shape_id} is not a chart")
    if any("points" in s for s in o["series"]):
        cd = XyChartData()
        for s in o["series"]:
            ser = cd.add_series(s["name"])
            for x, y in s["points"]:
                ser.add_data_point(x, y)
    else:
        # keep the chart's own value format ("$0.0M"…): axes linked to the source would otherwise fall back to General
        fmt = sh.chart._chartSpace.xpath(".//c:ser[1]/c:val//c:formatCode/text()")
        cd = CategoryChartData(number_format=o.get("number_format") or (fmt[0] if fmt else "General"))
        cd.categories = o["categories"]
        for s in o["series"]:
            if len(s["values"]) != len(o["categories"]):
                raise EditError(f"chart: series {s['name']!r} has {len(s['values'])} values for {len(o['categories'])} categories")
            cd.add_series(s["name"], s["values"])
    sh.chart.replace_data(cd)
    return f"chart     slide {o['slide']} #{sh.shape_id}: {len(o['series'])} series"


def op_table(d, o, w):
    slide = d.slide(o["slide"], "table")
    sh = find_shape(slide, o["shape"], "table")
    if not getattr(sh, "has_table", False) or not sh.has_table:
        raise EditError(f"table: #{sh.shape_id} is not a table")
    tbl = sh.table
    n = 0
    cells = o.get("cells")
    if cells is None:
        cells = [[None] * (o["col"]) + [o["text"]]]
        r0 = o["row"]
        cells = [None] * r0 + cells
    for r, row in enumerate(cells):
        if row is None:
            continue
        for c, val in enumerate(row):
            if val is None:
                continue
            if r >= len(tbl.rows) or c >= len(tbl.columns):
                raise EditError(f"table: cell ({r},{c}) is outside the {len(tbl.rows)}×{len(tbl.columns)} table "
                                "(adding rows/columns in place isn't supported — rebuild the slide via import)")
            set_text(tbl.cell(r, c).text_frame._txBody, val)
            n += 1
    return f"table     slide {o['slide']} #{sh.shape_id}: {n} cell(s)"


def op_delete_slide(d, o, w):
    slide = d.slide(o["slide"], "delete_slide")
    delete_slide(d, slide)
    return f"delete    slide {o['slide']}"


def op_move_slide(d, o, w):
    slide = d.slide(o["slide"], "move_slide")
    d.place(slide, "move_slide", o.get("after"), o.get("before"), o.get("at"))
    return f"move      slide {o['slide']} -> " + next(f"{k} {o[k]}" for k in ("after", "before", "at") if k in o)


def op_hide_slide(d, o, w):
    slide = d.slide(o["slide"], "hide_slide")
    if o.get("hidden", True):
        slide._element.set("show", "0")
    elif "show" in slide._element.attrib:
        del slide._element.attrib["show"]
    return f"hide      slide {o['slide']} = {o.get('hidden', True)}"


def _labels(o, n):
    labels = o.get("as")
    if labels is None:
        return [None] * n
    labels = [labels] if isinstance(labels, str) else list(labels)
    if len(labels) != n:
        raise EditError(f"{o['op']}: {n} slide(s) but {len(labels)} label(s) in \"as\"")
    return labels


def _register(d, labels, slides, op):
    for lab, s in zip(labels, slides):
        if lab is None:
            continue
        if lab in d.refs:
            raise EditError(f"{op}: label {lab!r} is already in use")
        d.refs[lab] = s


def _insert_block(d, new, o, op, default_after):
    """Place newly appended slides as one block according to after/before/at/replace."""
    if o.get("replace") is not None:
        target = d.slide(o["replace"], op)
        anchor_after = None
        pos = d.position(target)
        delete_slide(d, target)
        for j, s in enumerate(new):
            d.place(s, op, at=pos + 1 + j)
        return anchor_after
    if not any(k in o for k in ("after", "before", "at")):
        if default_after is None:
            return
        o = dict(o, after=default_after)
    first = new[0]
    d.place(first, op, o.get("after"), o.get("before"), o.get("at"))
    for prev, s in zip(new, new[1:]):
        lst = d.prs.slides._sldIdLst
        el = next(x for x in lst if int(x.get("id")) == s.slide_id)
        pel = next(x for x in lst if int(x.get("id")) == prev.slide_id)
        lst.remove(el)
        pel.addnext(el)


def op_import(d, o, w):
    from .compile import open_template
    src_path = pathlib.Path(o["from"])
    if not src_path.exists():
        raise EditError(f"import: {src_path} not found")
    src = open_template(src_path)
    picks = o.get("slides") or list(range(1, len(src.slides) + 1))
    src_slides = list(src.slides)
    for n in picks:
        if not 1 <= int(n) <= len(src_slides):
            raise EditError(f"import: {src_path.name} has no slide {n}")
    labels = _labels(o, len(picks))
    memo = {"parts": {}, "reserved": set()}
    new = [clone_slide(d.prs, src_slides[int(n) - 1], o.get("layout", "blank"), memo) for n in picks]
    if (src.slide_width, src.slide_height) != (d.prs.slide_width, d.prs.slide_height):
        w.append(f"import: {src_path.name} has a different slide size; imported content is not rescaled")
    _register(d, labels, new, "import")
    _insert_block(d, new, o, "import", None)
    d._keep_alive = getattr(d, "_keep_alive", []) + [src]
    return f"import    {len(new)} slide(s) from {src_path.name}" + (f" replacing slide {o['replace']}" if o.get("replace") is not None else "")


def op_duplicate(d, o, w):
    slide = d.slide(o["slide"], "duplicate")
    labels = _labels(o, 1)
    new = clone_slide(d.prs, slide, "match")
    _register(d, labels, [new], "duplicate")
    _insert_block(d, [new], o, "duplicate", o["slide"])
    return f"duplicate slide {o['slide']}" + (f" as {labels[0]!r}" if labels[0] else "")


OPS = {"text": op_text, "replace": op_replace, "notes": op_notes, "style": op_style, "geometry": op_geometry,
       "delete_shape": op_delete_shape, "image": op_image, "chart": op_chart, "table": op_table,
       "delete_slide": op_delete_slide, "move_slide": op_move_slide, "hide_slide": op_hide_slide,
       "import": op_import, "duplicate": op_duplicate}


def apply_edits(src, edits, out, backup=False, author=""):
    from .compile import open_template, reset_core_props
    src, out = pathlib.Path(src), pathlib.Path(out)
    spec = read_json(edits) if isinstance(edits, (str, pathlib.Path)) else edits
    ops = spec["edits"] if isinstance(spec, dict) else spec
    if src.suffix.lower() not in (".pptx", ".potx"):
        raise EditError(f"edit works on .pptx files; convert {src.suffix} first (inspect converts .ppt/.odp)")
    prs = open_template(src)
    d = Deck(prs)
    membership = _section_membership(prs)
    log, warnings = [], []
    base = pathlib.Path(edits).parent if isinstance(edits, (str, pathlib.Path)) else pathlib.Path.cwd()
    for n, o in enumerate(ops, 1):
        fn = OPS.get(o.get("op"))
        if fn is None:
            raise EditError(f"edit {n}: unknown op {o.get('op')!r}. Known: {', '.join(OPS)}")
        for key in ("from", "src"):
            if key in o and not pathlib.Path(o[key]).is_absolute():
                o = dict(o, **{key: str((base / o[key]).resolve())})
        try:
            log.append(fn(d, o, warnings))
        except EditError as e:
            raise EditError(f"edit {n} ({o.get('op')}): {e} — nothing was written") from None
        except KeyError as e:
            raise EditError(f"edit {n} ({o.get('op')}): missing field {e} — nothing was written") from None
    sync_sections(prs, membership)
    # add_slide names parts slide<count+1>.xml, which collides after a delete: renumber before saving
    prs.part.rename_slide_parts([sid.rId for sid in prs.slides._sldIdLst])
    if backup and out.resolve() == src.resolve():
        shutil.copy2(src, src.with_suffix(src.suffix + ".bak"))
    # their deck: keep its author, only stamp who touched it last
    reset_core_props(prs, author, keep_author=True)
    prs.save(out)
    return {"log": log, "warnings": warnings, "output": str(out), "slides": len(prs.slides)}
