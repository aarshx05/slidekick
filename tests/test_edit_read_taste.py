"""Edit mode keeps formatting, Analyse mode reads decks back, and the taste lint catches what it claims to."""
import json

import pytest
from pptx import Presentation


def _first_text_shape(pptx, needle):
    prs = Presentation(pptx)
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.has_text_frame and needle in sh.text_frame.text:
                return i, sh
    raise AssertionError(f"no shape containing {needle!r}")


def test_edit_text_keeps_run_formatting(tmp_path, built):
    from deckkit.edit import apply_edits

    n, sh = _first_text_shape(built["pptx"], "writes itself")
    before = sh.text_frame.paragraphs[0].runs[0].font
    size, bold = before.size, before.bold
    edits = tmp_path / "edits.json"
    edits.write_text(json.dumps({"edits": [{"op": "text", "slide": n, "shape": sh.shape_id, "text": "Decks that build themselves"}]}))
    out = tmp_path / "edited.pptx"
    r = apply_edits(built["pptx"], edits, out, author="Reviewer")
    assert r["warnings"] == []
    _, after = _first_text_shape(out, "Decks that build themselves")
    run = after.text_frame.paragraphs[0].runs[0].font
    assert (run.size, run.bold) == (size, bold)
    cp = Presentation(out).core_properties
    assert cp.last_modified_by == "Reviewer"


def test_edit_rejects_unknown_op_without_writing(tmp_path, built):
    from deckkit.edit import EditError, apply_edits

    edits = tmp_path / "bad.json"
    edits.write_text(json.dumps({"edits": [{"op": "explode"}]}))
    out = tmp_path / "never.pptx"
    with pytest.raises(EditError):
        apply_edits(built["pptx"], edits, out)
    assert not out.exists()


def test_inspect_reads_outline(tmp_path, built):
    from deckkit.reader import inspect_deck

    inv = inspect_deck(built["pptx"], tmp_path / "inspect", render=False)
    assert len(inv["slides"]) == len(built["layout"]["slides"])
    assert (tmp_path / "inspect" / "inventory.json").exists()


def test_taste_flags_starter_kicker_habit(built):
    from deckkit.taste import run_taste

    r = run_taste(project=built["project"], out_dir=built["project"] / "qa")
    codes = {f["code"] for f in r["findings"]}
    assert "kicker" in codes  # the starter puts a kicker on nearly every slide on purpose
    assert (built["project"] / "qa" / "taste.md").exists()
