"""Regression tests for issues found while building the gallery decks."""
import json
import pathlib

from conftest import BRAND_KIT
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


def _paths(node):
    if isinstance(node, dict):
        for v in node.values():
            yield from _paths(v)
    elif isinstance(node, list):
        for v in node:
            yield from _paths(v)
    elif isinstance(node, str) and ("/" in node or "\\" in node):
        yield node


def test_brand_json_paths_are_relative_and_resolve(tmp_path, monkeypatch):
    import shutil

    from deckkit.brand import ingest
    # the kit sits beside the project (same drive: CI keeps the repo on D: and temp on C:), passed as a relative path,
    # which used to be stored relative to the cwd and then read relative to brand/
    shutil.copytree(BRAND_KIT, tmp_path / "kit")
    monkeypatch.chdir(tmp_path)
    out = tmp_path / "proj" / "brand"
    ingest(["kit"], out)
    data = json.loads((out / "brand.json").read_text(encoding="utf-8"))
    for p in _paths(data):
        assert not pathlib.Path(p).is_absolute(), f"absolute path leaked into brand.json: {p}"
    assert (out / data["template"]).exists()
    for logo in data["logos"]:
        assert (out / logo["path"]).exists()


def test_chart_edit_keeps_number_format(tmp_path):
    from deckkit.edit import apply_edits

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    cd = CategoryChartData(number_format='"$"0.0"M"')
    cd.categories = ["Q1", "Q2"]
    cd.add_series("Revenue", (14.8, 16.9))
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(6), Inches(4), cd)
    src = tmp_path / "chart.pptx"
    prs.save(src)
    edits = tmp_path / "edits.json"
    edits.write_text(json.dumps({"edits": [{"op": "chart", "slide": 1, "shape": gf.shape_id,
                                            "categories": ["Q3", "Q4"], "series": [{"name": "Revenue", "values": [18.4, 19.8]}]}]}))
    out = tmp_path / "out.pptx"
    apply_edits(src, edits, out)
    chart = next(sh for sh in Presentation(out).slides[0].shapes if sh.has_chart).chart
    assert chart._chartSpace.xpath(".//c:ser[1]/c:val//c:formatCode/text()") == ['"$"0.0"M"']
    assert list(chart.plots[0].categories) == ["Q3", "Q4"]


def test_fetch_and_fonts_follow_project(tmp_path):
    import argparse

    from deckkit import cli
    seen = {}
    import deckkit.capture as cap
    orig = cap.fetch
    cap.fetch = lambda url, out, **kw: seen.setdefault("out", out)
    try:
        cli.cmd_fetch(argparse.Namespace(url="https://example.com/x.png", out=None, project=str(tmp_path), name=None, note=None))
    finally:
        cap.fetch = orig
    assert pathlib.Path(seen["out"]) == tmp_path.resolve() / "assets"
