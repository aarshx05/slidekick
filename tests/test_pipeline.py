"""Canvas -> layout.json -> .pptx: the deck that comes out is native, editable and carries no stray identity."""
import zipfile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def test_extract_is_clean(built):
    layout = built["layout"]
    assert len(layout["slides"]) >= 5
    assert layout["warnings"] == [], layout["warnings"]


def test_every_slide_compiles_to_native_objects(built):
    prs = Presentation(built["pptx"])
    assert len(prs.slides) == len(built["layout"]["slides"])
    for slide in prs.slides:
        assert len(slide.shapes) > 0
    texts = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    assert any("writes itself" in t for t in texts), "headline should be real, editable text"


def test_chart_has_embedded_workbook(built):
    prs = Presentation(built["pptx"])
    charts = [sh for s in prs.slides for sh in s.shapes if sh.has_chart]
    if not charts:
        return
    assert charts[0].chart.part.chart_workbook.xlsx_part is not None


def test_groups_survive(built):
    prs = Presentation(built["pptx"])
    kinds = {sh.shape_type for s in prs.slides for sh in s.shapes}
    assert MSO_SHAPE_TYPE.GROUP in kinds


def test_transitions_written(built):
    with zipfile.ZipFile(built["pptx"]) as z:
        xml = b"".join(z.read(n) for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml"))
    assert b"<p:transition" in xml
    assert b"morph" in xml


def test_core_properties_are_scrubbed(built):
    cp = Presentation(built["pptx"]).core_properties
    assert cp.author == ""
    assert cp.last_modified_by == ""  # python-pptx's default template says "Steve Canny"


def test_template_author_does_not_leak(tmp_path, built):
    from conftest import BRAND_KIT
    from deckkit.compile import open_template, reset_core_props

    prs = open_template(BRAND_KIT / "deck-template.potx")
    prs.core_properties.author = "Someone Else"
    reset_core_props(prs, "Deck Bot")
    out = tmp_path / "t.pptx"
    prs.save(out)
    cp = Presentation(out).core_properties
    assert (cp.author, cp.last_modified_by, cp.revision) == ("Deck Bot", "Deck Bot", 1)


def test_chart_point_colors_and_labels(tmp_path):
    """One highlighted bar is a native per-point colour, not a fake extra series."""
    import argparse
    import json

    from deckkit.cli import cmd_init
    from deckkit.compile import compile_deck
    from deckkit.extract import extract

    project = tmp_path / "p"
    cmd_init(argparse.Namespace(dir=str(project), force=False))
    spec = {"type": "column", "categories": ["Q1", "Q2", "Q3", "Q4"], "series": [{"name": "Revenue", "values": [12, 18, 27, 31],
            "color": "#BBBBBB", "point_colors": {"3": "#F26B1D"}}], "labels": True, "label_position": "inside_end", "label_color": "#FFFFFF"}
    (project / "canvas" / "deck.html").write_text(
        '<!doctype html><html><head><link rel="stylesheet" href="deck.css"><link rel="stylesheet" href="brand.css"></head><body>'
        '<section class="slide"><h2 data-pptx="text" class="abs" style="left:64px;top:64px;width:900px;font:700 40px/48px sans-serif">'
        'Revenue hit a record in Q4</h2><div data-pptx="chart" class="abs" style="left:64px;top:160px;width:800px;height:420px;'
        f"font:400 16px/20px sans-serif\" data-chart='{json.dumps(spec)}'></div></section></body></html>", encoding="utf-8")
    extract(project)
    prs = Presentation(compile_deck(project, calibrate=False)["output"])
    chart = next(sh for sh in prs.slides[0].shapes if sh.has_chart).chart
    xml = chart._chartSpace.xml
    assert len(chart.plots[0].series) == 1
    assert "<c:dPt>" in xml and 'val="F26B1D"' in xml and 'val="BBBBBB"' in xml
    assert '<c:dLblPos val="inEnd"/>' in xml
